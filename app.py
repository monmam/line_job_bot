import os
import re
import json
import copy
import time
import threading
import datetime

import requests
from flask import Flask, request, abort, render_template, jsonify, send_file
from dotenv import load_dotenv

# Import LINE SDK (v3)
from linebot.v3 import WebhookHandler
from linebot.v3.exceptions import InvalidSignatureError
from linebot.v3.messaging import (
    Configuration, ApiClient, MessagingApi, PushMessageRequest, ReplyMessageRequest, TextMessage
)
from linebot.v3.webhooks import MessageEvent, TextMessageContent, JoinEvent, LeaveEvent

load_dotenv()

# โหลด .env ก่อน แล้วค่อย import โมดูลที่อ่านค่า env
import storage
from google_sheet import (
    save_to_google_sheets, get_sheet_client, delete_row_from_google_sheets, has_credentials
)

app = Flask(__name__)

LINE_CHANNEL_ACCESS_TOKEN = os.getenv("LINE_CHANNEL_ACCESS_TOKEN", "")
LINE_CHANNEL_SECRET = os.getenv("LINE_CHANNEL_SECRET", "")
GOOGLE_SPREADSHEET_ID = os.getenv("GOOGLE_SPREADSHEET_ID", "")

handler = WebhookHandler(LINE_CHANNEL_SECRET)
configuration = Configuration(access_token=LINE_CHANNEL_ACCESS_TOKEN)

# บอทสำรอง (ส่งอย่างเดียว) ใช้เมื่อบอทหลักส่ง push ไม่ได้ เช่น โควต้ารายเดือนเต็ม
# ใส่ใน Render: LINE_BACKUP_TOKENS=tokenบอท2,tokenบอท3
BACKUP_CONFIGURATIONS = [
    Configuration(access_token=t.strip())
    for t in os.getenv("LINE_BACKUP_TOKENS", "").split(",") if t.strip()
]


def bots_status():
    """สถานะบอททั้งหมด: บอทหลัก + บอทสำรองที่ใส่ใน LINE_BACKUP_TOKENS (รับข้อความได้เมื่อมี secret ตรงลำดับ)"""
    bots = [{"name": "บอทหลัก", "receiving": bool(LINE_CHANNEL_SECRET)}]
    for i in range(len(BACKUP_CONFIGURATIONS)):
        bots.append({"name": f"บอทสำรอง {i + 1}",
                     "receiving": i < len(_backup_secrets) and bool(_backup_secrets[i])})
    return bots


# ตอบกลับต้องใช้ access token ของ "บอทตัวที่ส่งข้อความมา" ถ้า LINE_BACKUP_TOKENS กับ LINE_BACKUP_SECRETS เรียงลำดับไม่ตรงกัน
# reply จะล้มด้วย "Invalid reply token" ระบบจึงถาม LINE ว่า token แต่ละตัวเป็นบอทไหน (userId) แล้วจับคู่กับ destination ใน webhook ให้เอง
_ctx = threading.local()
_bot_ids = {}                 # userId ของบอท -> (ชื่อ, Configuration)
_bot_meta = {}                # key บอท (main / backup1 ...) -> {display, basic}
_bot_ids_loaded = False
_bot_ids_lock = threading.Lock()
_pair_warned = set()


def _load_bot_ids():
    global _bot_ids_loaded
    with _bot_ids_lock:
        if _bot_ids_loaded:
            return
        _bot_ids_loaded = True
        names = ["บอทหลัก"] + [f"บอทสำรอง {i + 1}" for i in range(len(BACKUP_CONFIGURATIONS))]
        keys = ["main"] + [f"backup{i + 1}" for i in range(len(BACKUP_CONFIGURATIONS))]
        for key, name, conf in zip(keys, names, [configuration] + BACKUP_CONFIGURATIONS):
            try:
                r = requests.get("https://api.line.me/v2/bot/info",
                                 headers={"Authorization": f"Bearer {conf.access_token}"}, timeout=8)
                if r.status_code == 200 and r.json().get("userId"):
                    _bot_ids[r.json()["userId"]] = (name, conf)
                    _bot_meta[key] = {"display": r.json().get("displayName", ""), "basic": r.json().get("basicId", "")}
                else:
                    print(f"⚠️ {name}: เช็ค access token ไม่ผ่าน (LINE ตอบ {r.status_code}) ใส่ Channel access token ถูกตัวหรือยัง")
            except Exception as e:
                print(f"⚠️ {name}: ตรวจ access token ไม่สำเร็จ: {e}")


def _pick_conf(default_conf, default_name):
    """เลือก token ที่ตรงกับบอทที่ส่ง webhook มาจริง ๆ (ดูจาก destination) ถ้าหาไม่ได้ใช้ตัวเดิม"""
    dest = getattr(_ctx, "destination", None)
    if not dest:
        return default_conf
    _load_bot_ids()
    hit = _bot_ids.get(dest)
    if not hit:
        return default_conf
    name, real_conf = hit
    if real_conf is not default_conf and (default_name, name) not in _pair_warned:
        _pair_warned.add((default_name, name))
        print(f"⚠️ {default_name}: Channel secret กับ access token ไม่ใช่บอทตัวเดียวกัน (บอทที่ส่งข้อความมาคือ {name}) "
              f"ใช้ token ของ {name} ตอบแทนให้แล้ว แต่ควรแก้ลำดับ LINE_BACKUP_SECRETS ให้ตรงกับ LINE_BACKUP_TOKENS")
    return real_conf


# ====================== สิทธิ์คำสั่ง: เปิด/ปิดได้ต่อบอท และต่อกลุ่ม (ตั้งค่าในหน้าเว็บ "บอท & คำสั่ง") ======================
# ใช้ได้ต่อเมื่อ "บอทเปิด" และ "กลุ่มเปิด" ทั้งคู่ (ไม่ได้ตั้งค่าไว้ = เปิดทั้งหมด)
COMMANDS = [
    {"id": "ticket", "name": "ย่อใบงาน", "desc": "โยนใบงานเต็มเข้ามา บอทอ่านแล้วย่อตอบกลับ", "example": "(วางใบงานเต็ม)"},
    {"id": "update", "name": "ดูตารางสรุป (Update)", "desc": "ตารางสรุปย่อของวันที่เลือก", "example": "Update 18"},
    {"id": "full_table", "name": "ดูตารางเต็ม", "desc": "ใบสรุปเต็มทุกงานของวัน (ไอคอนก่อนใส่ทะเบียน)", "example": "ดูตารางเต็ม 18"},
    {"id": "view_job", "name": "ดูงานเดียว", "desc": "ใบสรุปเต็มของงานเดียว", "example": "ดูงาน F01-18"},
    {"id": "cut_table", "name": "ตัดตาราง", "desc": "ใบสรุปเต็มเฉพาะงานที่ยังไม่มีทะเบียน", "example": "ตัดตาราง 18"},
    {"id": "now", "name": "งานด่วน NOW", "desc": "ข้อความงานด่วน: ขนาดรถ / เส้นทาง / ราคา", "example": "now F01"},
    {"id": "close_job", "name": "ปิดใบงาน (ใส่ทะเบียน)", "desc": "ข้อมูลคนขับ + ใบงาน ใส่ทะเบียนแล้วตอบข้อความรวม", "example": "CAR NO : ... + ใบงาน"},
    {"id": "return", "name": "คืนงาน (📌)", "desc": "เปลี่ยนป้ายงานเป็น 📌", "example": "F22 คืน 16"},
    {"id": "cancel", "name": "ยกเลิกงาน", "desc": "ใบงาน + บรรทัด cancel", "example": "ใบงาน + cancel"},
    {"id": "uncancel", "name": "ดึงงานที่ยกเลิกกลับ", "desc": "ยกเลิกการยกเลิก", "example": "uncancel F09 / un F09"},
    {"id": "price", "name": "แก้ไขราคา", "desc": "ปรับราคาต่อท้ายในตาราง", "example": "แก้ไขราคา / F01 +50"},
    {"id": "edit_job", "name": "แก้ไขใบสรุป", "desc": "แก้เวลา / เส้นทาง / ขนาดรถ / Order ในใบสรุป (ค่าเดิม>ค่าใหม่)", "example": "VK01 19 13:25>20:00"},
    {"id": "clear_plate", "name": "ล้างทะเบียน", "desc": "เอาทะเบียนหรือ 📌 ออก", "example": "ล้างทะเบียน VK05"},
    {"id": "group_id", "name": "ขอ Group ID", "desc": "บอทตอบรหัสกลุ่ม (พิมพ์ในกลุ่ม)", "example": "id"},
    {"id": "help", "name": "เมนูคำสั่ง", "desc": "แสดงรายการคำสั่งทั้งหมด", "example": "help"},
]
COMMAND_IDS = [c["id"] for c in COMMANDS]
_seen_groups = {}          # group_id -> {"bot": key, "name": ชื่อกลุ่ม, "last": เวลา}
_seen_lock = threading.Lock()


def _bot_list():
    """[(key, ชื่อ, Configuration)] บอทหลัก + บอทสำรอง"""
    return [("main", "บอทหลัก", configuration)] + [
        (f"backup{i + 1}", f"บอทสำรอง {i + 1}", conf) for i, conf in enumerate(BACKUP_CONFIGURATIONS)]


def _bot_key_of(conf):
    for key, _, c in _bot_list():
        if c is conf:
            return key
    return "main"


def _cmd_allowed(cmd, bot_key, group_id=None):
    st = load_settings()
    if (st.get("bot_commands", {}).get(bot_key, {}) or {}).get(cmd, True) is False:
        return False
    if group_id and (st.get("group_commands", {}).get(group_id, {}) or {}).get(cmd, True) is False:
        return False
    return True


def _record_group(group_id, bot_key):
    """จำกลุ่มที่บอทอยู่ (ไว้แสดงในหน้าตั้งค่าสิทธิ์) เก็บใน Supabase"""
    if not group_id:
        return
    with _seen_lock:
        entry = _seen_groups.setdefault(group_id, {})
        changed = entry.get("bot") != bot_key
        entry["bot"] = bot_key
        entry["last"] = _ts()
        snapshot = json.loads(json.dumps(_seen_groups)) if changed else None
    if snapshot is not None:
        try:
            storage.set("seen_groups", snapshot)
        except Exception as e:
            print(f"⚠️ บันทึกรายชื่อกลุ่มไม่สำเร็จ: {e}")


def _fetch_group_name(group_id, conf):
    if not str(group_id).startswith("C"):
        return ""
    try:
        r = requests.get(f"https://api.line.me/v2/bot/group/{group_id}/summary",
                         headers={"Authorization": f"Bearer {conf.access_token}"}, timeout=6)
        if r.status_code == 200:
            return r.json().get("groupName", "")
    except Exception:
        pass
    return ""


def _normalize_cmds(d):
    """เก็บเฉพาะคำสั่งที่ปิด (False) และเป็นรหัสคำสั่งที่มีจริง"""
    return {c: False for c in COMMAND_IDS if isinstance(d, dict) and d.get(c) is False}


@app.route("/api/permissions", methods=["GET"])
def api_permissions():
    _load_bot_ids()
    st = load_settings()
    bc, gc = st.get("bot_commands", {}), st.get("group_commands", {})
    full = lambda disabled: {c: (disabled or {}).get(c, True) is not False for c in COMMAND_IDS}
    names = {}
    bots = []
    for key, name, conf in _bot_list():
        meta = _bot_meta.get(key, {})
        names[key] = name
        bots.append({"key": key, "name": name, "display": meta.get("display", ""), "basic": meta.get("basic", ""),
                     "receiving": any(b["name"] == name and b["receiving"] for b in bots_status()),
                     "commands": full(bc.get(key))})
    listed = {str(g.get("group_id", "")).strip(): g.get("name", "") for g in st.get("line_groups", []) if g.get("group_id")}
    with _seen_lock:
        seen = json.loads(json.dumps(_seen_groups))
    ids = list(dict.fromkeys(list(seen) + list(listed) + list(gc)))
    fetched, groups = 0, []
    conf_of = {k: c for k, _, c in _bot_list()}
    for gid in ids:
        info = seen.get(gid, {})
        name = listed.get(gid) or info.get("name") or ""
        if not name and fetched < 8 and info.get("bot") in conf_of:
            fetched += 1
            name = _fetch_group_name(gid, conf_of[info["bot"]])
            if name:
                with _seen_lock:
                    _seen_groups.setdefault(gid, {})["name"] = name
                try:
                    storage.set("seen_groups", json.loads(json.dumps(_seen_groups)))
                except Exception:
                    pass
        groups.append({"id": gid, "name": name or gid, "bot": info.get("bot", ""), "bot_name": names.get(info.get("bot", ""), ""),
                       "last": info.get("last", ""), "commands": full(gc.get(gid))})
    return jsonify({"commands": COMMANDS, "bots": bots, "groups": groups})


@app.route("/api/save_permissions", methods=["POST"])
def api_save_permissions():
    data = request.get_json(silent=True) or {}
    bot_commands = {k: _normalize_cmds(v) for k, v in (data.get("bot_commands") or {}).items()}
    group_commands = {k: _normalize_cmds(v) for k, v in (data.get("group_commands") or {}).items()}
    ok, msg = save_settings({"bot_commands": {k: v for k, v in bot_commands.items() if v},
                             "group_commands": {k: v for k, v in group_commands.items() if v}})
    return jsonify({"success": ok, "message": "บันทึกสิทธิ์คำสั่งเรียบร้อยแล้ว" if ok else msg}), (200 if ok else 500)


def push_with_fallback(to, messages, label):
    """ส่ง push ด้วยบอทหลัก ถ้าไม่ได้ (เช่น โควต้าเต็ม / บอทไม่อยู่ในกลุ่ม) ไล่ใช้บอทสำรองทีละตัว"""
    for idx, conf in enumerate([configuration] + BACKUP_CONFIGURATIONS):
        bot_name = "บอทหลัก" if idx == 0 else f"บอทสำรอง {idx}"
        try:
            with ApiClient(conf) as api_client:
                MessagingApi(api_client).push_message(PushMessageRequest(to=to, messages=messages))
            print(f"✅ ส่งไปยัง {label} ด้วย {bot_name} สำเร็จ")
            return True
        except Exception as e:
            status = getattr(e, "status", "")
            body = str(getattr(e, "body", "") or e)[:200]
            print(f"❌ ส่งไปยัง {label} ด้วย {bot_name} ล้มเหลว ({status}): {body}")
    return False


def _reply_text(conf, reply_token, text):
    try:
        with ApiClient(conf) as api_client:
            MessagingApi(api_client).reply_message(
                ReplyMessageRequest(reply_token=reply_token, messages=[TextMessage(text=text)]))
    except Exception as e:
        print(f"⚠️ ตอบกลับไม่ได้: {e}")


# บอทสำรองรับคำสั่ง "id" ในกลุ่มได้ด้วย (ตอบ Group ID) ต้องตั้ง Webhook URL ของบอทสำรองเป็น /callback เหมือนบอทหลัก
# และใส่ Channel secret ของบอทสำรองใน Render: LINE_BACKUP_SECRETS=secretบอท2,secretบอท3 (เรียงตรงกับ LINE_BACKUP_TOKENS)
def _make_backup_handler(secret, conf, name):
    h = WebhookHandler(secret)

    @h.add(MessageEvent, message=TextMessageContent)
    def _on_message(event):
        print(f"🤖 ข้อความเข้าบอท {name}")
        _handle_text(event, _pick_conf(conf, name))

    @h.add(JoinEvent)
    def _on_join(event):
        src = event.source
        gid = getattr(src, 'group_id', None) or getattr(src, 'room_id', None)
        print(f"➕ {name} ถูกเชิญเข้า {src.type}: {gid}")
        _record_group(gid, _bot_key_of(conf))

    @h.add(LeaveEvent)
    def _on_leave(event):
        src = event.source
        print(f"➖ {name} ถูกนำออกจาก {src.type}: {getattr(src, 'group_id', None) or getattr(src, 'room_id', None)}")

    return h


_backup_secrets = [x.strip() for x in os.getenv("LINE_BACKUP_SECRETS", "").split(",")]
BACKUP_HANDLERS = [
    _make_backup_handler(_backup_secrets[i], conf, f"บอทสำรอง {i + 1}")
    for i, conf in enumerate(BACKUP_CONFIGURATIONS)
    if i < len(_backup_secrets) and _backup_secrets[i]
]

# เวลาไทย (เซิร์ฟเวอร์ Render ใช้เวลา UTC)
TH_TZ = datetime.timezone(datetime.timedelta(hours=7))


def now_th():
    return datetime.datetime.now(TH_TZ)


# Global State (โหลดจาก Supabase ตอนเปิดเครื่อง และบันทึกกลับทุกครั้งที่เปลี่ยน)
state_lock = threading.RLock()
job_queue = []
timer_thread = None
timer_start_time = None
latest_jobs = []
last_sender_id = None
last_reply = {}  # chat_id -> {"token", "conf"} reply token ล่าสุดของแต่ละแชท ใช้ตอบกลับแชทนั้นฟรี (ไม่นับโควต้า push)
# PUSH_ENABLED=1 = ส่งใบสรุปเข้ากลุ่ม/ผู้ส่งด้วย push แบบเดิม (เสียโควต้า) ค่าเริ่มต้น 0 = ตอบกลับ (reply) แชทที่พิมพ์มาเท่านั้น
PUSH_ENABLED = os.getenv("PUSH_ENABLED", "0").strip() == "1"
pending_update_dates = []  # วันที่ที่มีการอัปเดตป้ายทะเบียน รอส่งใบสรุปของวันนั้นใหม่
MAX_HISTORY = 1000     # จำนวนใบงานสูงสุดที่เก็บในประวัติ
KEEP_DAYS = 14         # เก็บใบงานย้อนหลังกี่วัน (ใบงานที่วันที่เก่ากว่านี้จะถูกตัดออก)
LOG_CAP = {"activity": 1000, "sent": 300}
_activity_log = []     # บันทึกเหตุการณ์: รับใบงาน / แก้ไข / ซ้ำ / ป้ายทะเบียน / คืน / ยกเลิก
_sent_log = []         # บันทึกการส่ง: ใบสรุป / รายการย่อ พร้อมปลายทางและผลลัพธ์
_logs_dirty = False

# ค่าเริ่มต้นพจนานุกรมรายชื่อถนนหลัก (อังกฤษ = ไทย)
DEFAULT_MAIN_ROADS_DICT = {
    "New Petchaburi": "เพชรบุรีตัดใหม่",
    "Sukhumvit": "สุขุมวิท",
    "Phaholyothin": "พหลโยธิน",
    "Petchaburi": "เพชรบุรี",
    "Phetchaburi Rd": "เพชรบุรี",
    "Phetburi": "เพชรบุรี",
    "Phra Nakhon": "พระนคร",
    "Rama 1": "พระรามที่ 1",
    "Rama 2": "พระรามที่ 2",
    "Rama 3": "พระรามที่ 3",
    "Rama 4": "พระรามที่ 4",
    "Rama 5": "พระรามที่ 5",
    "Rama 6": "พระรามที่ 6",
    "Rama 7": "พระรามที่ 7",
    "Rama 8": "พระรามที่ 8",
    "Rama 9": "พระรามที่ 9",
    "Sathon": "สาทร",
    "Silom": "สีลม",
    "North Sathon Road": "สาทรเหนือ",
    "Charoennakorn Road": "เจริญนคร",
    "Surawong": "สุรวงศ์",
    "Ratchadamnoen Klang": "ราชดำเนินกลาง",
    "Ratchadamnoen Nok": "ราชดำเนินนอก",
    "Ratchadamnoen Nai": "ราชดำเนินใน",
    "Charoen Krung": "เจริญกรุง",
    "Bamrung Mueang": "บำรุงเมือง",
    "Din So": "ดินสอ",
    "Tanao": "ตะนาว",
    "Khaosan": "ข้าวสาร",
    "Phra Sumen": "พระสุเมรุ",
    "Samsen": "สามเสน",
    "Witthayu": "วิทยุ",
    "Lang Suan": "หลังสวน",
    "Chit Lom": "ชิดลม",
    "Ploenchit": "เพลินจิต",
    "Ratchadamri": "ราชดำริ",
    "Henri Dunant": "อังรีดูนังต์",
    "Phaya Thai": "พญาไท",
    "Banthat Thong": "บรรทัดทอง",
    "Chan": "จันทน์",
    "Sathu Pradit": "สาธุประดิษฐ์",
    "Nang Linchi": "นางลิ้นจี่",
    "Chuea Phloeng": "เชื้อเพลิง",
    "Naradhiwas Rajanagarindra": "นราธิวาสราชนครินทร์",
    "Ratchadaphisek": "รัชดาภิเษก",
    "Asok Montri": "อโศกมนตรี",
    "Thong Lo": "ทองหล่อ",
    "Ekkamai": "เอกมัย",
    "Pridi Banomyong": "ปรีดี พนมยงค์",
    "On Nut": "อ่อนนุช",
    "Bangna-Trat": "บางนา-ตราด",
    "Srinakarin": "ศรีนครินทร์",
    "Phatthanakan": "พัฒนาการ",
    "Ramkhamhaeng": "รามคำแหง",
    "Lat Phrao": "ลาดพร้าว",
    "Pradit Manutham": "ประดิษฐ์มนูธรรม",
    "Ram Inthra": "รามอินทรา",
    "Chaeng Watthana": "แจ้งวัฒนะ",
    "Ngam Wong Wan": "งามวงศ์วาน",
    "Tiwanon": "ติวานนท์",
    "Prachachuen": "ประชาชื่น",
    "Kamphaeng Phet": "กำแพงเพชร",
    "Vibhavadi Rangsit": "วิภาวดีรังสิต",
    "Sutthisan Winitchai": "สุทธิสารวินิจฉัย",
    "Pracha Uthit": "ประชาอุทิศ",
    "Phutthamonthon Sai 1": "พุทธมณฑลสาย 1",
    "Phutthamonthon Sai 2": "พุทธมณฑลสาย 2",
    "Phutthamonthon Sai 3": "พุทธมณฑลสาย 3",
    "Phutthamonthon Sai 4": "พุทธมณฑลสาย 4",
    "Borommaratchachonnani": "บรมราชชนนี",
    "Charan Sanitwong": "จรัญสนิทวงศ์",
    "Arun Amarin": "อรุณอมรินทร์",
    "Itsaraphap": "อิสรภาพ",
    "Prachathipok": "ประชาธิปก",
    "Somdej Phra Chao Tak Sin": "สมเด็จพระเจ้าตากสิน",
    "Charoen Nakhon": "เจริญนคร",
    "Rat Burana": "ราษฎร์บูรณะ",
    "Suksawat": "สุขสวัสดิ์",
    "Ekkachai": "เอกชัย",
    "Bang Khun Thian-Cha Thale": "บางขุนเทียน-ชายทะเล",
    "Kanchanaphisek": "กาญจนาภิเษก",
    "Outer Ring Road": "วงแหวนรอบนอก",
    "Phrannok": "พรานนก"
}

# ค่าเริ่มต้นพจนานุกรมรายชื่อย่านสำคัญ
DEFAULT_MAJOR_AREAS_DICT = {
    "Sukhumvit Road": "ถนนสุขุมวิท",
    "Khlong San": "คลองสาน",
    "Phahonyothin": "ถนนพหลโยธิน",
    "Phetchaburi": "ถนนเพชรบุรี",
    "Phahonyothin Road": "ถนนพหลโยธิน",
    "Phetchaburi Road": "ถนนเพชรบุรี",
    "Rama 1 Road": "ถนนพระรามที่ 1",
    "Rama 2 Road": "ถนนพระรามที่ 2",
    "Rama 3 Road": "ถนนพระรามที่ 3",
    "Rama 4 Road": "ถนนพระรามที่ 4",
    "Silom Road": "ถนนสีลม",
    "Sathorn Road": "ถนนสาทร",
    "Asok": "อโศก",
    "Thong Lo": "ทองหล่อ",
    "Ekkamai": "เอกมัย",
    "Siam": "สยาม",
    "Pratunam": "ประตูน้ำ",
    "Khaosan": "ข้าวสาร",
    "Silom": "สีลม",
    "Sathorn": "สาทร"
}

# Dictionary สำหรับ Map จุดรับ จุดส่ง ขนาดรถ และราคาตามกฎ
CAR_PRICING_MAP = {
    "5 SEAT": {"code": "5S", "price": "380"},
    "5S": {"code": "5S", "price": "380"},
    "7 SEAT": {"code": "7S", "price": "480"},
    "7S": {"code": "7S", "price": "480"},
    "CAM/7S": {"code": "Cam7S", "price": "480"},
    "CAMRY/7S": {"code": "Cam7S", "price": "480"},
    "CAMRY OR 7SEAT": {"code": "Cam7S", "price": "480"},
    "PREMIUM EV": {"code": "EV-SUV", "price": "450"},
    "EV-SUV": {"code": "EV-SUV", "price": "450"},
    "EV SUV": {"code": "EV-SUV", "price": "450"},
    "PREMIUM CAR": {"code": "Cam", "price": "480"},
    "CAMRY": {"code": "Cam", "price": "480"},
    "CAM": {"code": "Cam", "price": "480"},
    "MPV": {"code": "7S", "price": "480"},
    "EV": {"code": "EV", "price": "380"},
    "ECONOMY VAN": {"code": "Van", "price": "700"},
    "VAN": {"code": "Van", "price": "700"},
}

# โรงแรม/สถานที่ที่ใบงานไม่มีที่อยู่ให้ (ชื่อตัวเล็ก = ผลลัพธ์)
BUILTIN_PLACES = {
    "the quarter chao phraya": "เจริญนคร",
}


def match_car(text):
    """หาขนาดรถจากข้อความ (คำยาวก่อน เช่น PREMIUM EV ชนะ EV) คืน dict หรือ None"""
    upper = str(text or "").upper()
    for key in sorted(CAR_PRICING_MAP.keys(), key=len, reverse=True):
        if re.search(r"(?<![A-Z0-9])" + re.escape(key) + r"(?![A-Z0-9])", upper):
            return CAR_PRICING_MAP[key]
    return None


# ====================== การตั้งค่า (เก็บใน Supabase) ======================
_settings_cache = None
_settings_lock = threading.RLock()
_settings_retry_at = 0
_settings_error = ""


def _with_defaults(data):
    data = dict(data or {})
    data.setdefault("wait_seconds", 10)
    data.setdefault("line_groups", [])
    for grp in data["line_groups"]:
        if isinstance(grp, dict):
            grp.setdefault("send", True)   # ส่งใบสรุปเข้ากลุ่ม
            grp.setdefault("read", True)   # อ่านใบงานที่โพสต์ในกลุ่ม (ปิด = บอทไม่อ่านงานในกลุ่มนั้น)
    data.setdefault("custom_keywords", [])
    data.setdefault("bot_commands", {})     # key บอท -> {คำสั่ง: False} (เก็บเฉพาะที่ปิด)
    data.setdefault("group_commands", {})   # group_id -> {คำสั่ง: False}
    data.setdefault("main_roads_dict", copy.deepcopy(DEFAULT_MAIN_ROADS_DICT))
    data.setdefault("major_areas_dict", copy.deepcopy(DEFAULT_MAJOR_AREAS_DICT))
    data.setdefault("custom_locations", {"metropole": "เพชรบุรีตัดใหม่", "c u inn": "จตุจักร"})
    data.pop("waiting_time", None)
    try:
        data["wait_seconds"] = int(data["wait_seconds"])
    except Exception:
        data["wait_seconds"] = 10
    return data


def load_settings(force=False):
    """อ่านการตั้งค่า (เก็บไว้ในหน่วยความจำ ไม่ต้องดึงจาก Supabase ทุกครั้ง)"""
    global _settings_cache, _settings_retry_at, _settings_error
    with _settings_lock:
        if _settings_cache is not None and not force:
            return _settings_cache
        if not force and time.time() < _settings_retry_at:
            return _with_defaults({})
        try:
            data = storage.get("settings")
            if data is None:
                # ครั้งแรก: ย้ายค่าจาก settings.json เดิม (ถ้ามี) ขึ้น Supabase
                data = _with_defaults(storage._read_legacy_settings() or {})
                try:
                    storage.set("settings", data)
                except storage.StorageError as e:
                    print(f"⚠️ บันทึกการตั้งค่าเริ่มต้นไม่สำเร็จ: {e}")
            _settings_cache = _with_defaults(data)
            _settings_error = ""
            return _settings_cache
        except storage.StorageError as e:
            _settings_error = str(e)
            _settings_retry_at = time.time() + 30
            print(f"❌ โหลดการตั้งค่าไม่สำเร็จ (ใช้ค่าเริ่มต้นชั่วคราว): {e}")
            return _with_defaults({})


def save_settings(data):
    """บันทึกการตั้งค่า คืน (สำเร็จไหม, ข้อความ)"""
    global _settings_cache
    with _settings_lock:
        current = copy.deepcopy(load_settings(force=True))
        if _settings_error:
            return False, f"บันทึกไม่สำเร็จ: เชื่อมต่อที่เก็บข้อมูลไม่ได้ ({_settings_error})"

        if "waiting_time" in data:
            try:
                val = float(data["waiting_time"])
                current["wait_seconds"] = int(val if val >= 5 else val * 60)
            except Exception:
                pass

        for key in ("line_groups", "custom_keywords", "main_roads_dict",
                    "major_areas_dict", "custom_locations", "bot_commands", "group_commands"):
            if key in data:
                current[key] = data[key]

        try:
            storage.set("settings", current)
        except storage.StorageError as e:
            return False, f"บันทึกไม่สำเร็จ: {e}"
        _settings_cache = current
        return True, "บันทึกการตั้งค่าเรียบร้อยแล้ว"


# ====================== บันทึก/โหลด คิวใบงาน ======================

# ====================== ระบบเก็บข้อมูลใบงาน (ใบเต็ม / ใบย่อ / คืน / ยกเลิก / ประวัติ) ======================
# ใบงานแต่ละใบ (job) เก็บ: ข้อความเต็ม (text) สรุปปัจจุบัน (formatted_summary) สถานะส่ง (status)
#   state = normal | returned (📌) | cancelled (❌CANCEL❌), plate, events (ประวัติของใบงานนั้น)
# ประวัติรวมเก็บแยก 2 ชุด: activity_log (เหตุการณ์) และ sent_log (การส่ง) ทั้งหมดบันทึกลง Supabase
def _ts():
    return now_th().strftime("%Y-%m-%d %H:%M:%S")


def log_activity(kind, job=None, detail=""):
    global _logs_dirty
    job = job or {}
    with state_lock:
        _activity_log.append({
            "ts": _ts(), "type": kind, "id": job.get("id", ""),
            "date": _norm_job_date(job.get("date")) if job.get("date") else "",
            "order": job.get("order", ""), "detail": detail,
        })
        del _activity_log[:-LOG_CAP["activity"]]
        _logs_dirty = True


def log_sent(kind, text, dates, results):
    """results = [(ปลายทาง, สำเร็จไหม, reply/push)]"""
    global _logs_dirty
    with state_lock:
        _sent_log.append({
            "ts": _ts(), "kind": kind, "text": text, "dates": list(dates),
            "targets": [{"to": t, "ok": bool(ok), "via": via} for t, ok, via in results],
        })
        del _sent_log[:-LOG_CAP["sent"]]
        _logs_dirty = True


def mark_job(job, kind, detail=""):
    """บันทึกเหตุการณ์ลงใบงาน + ประวัติรวม kind: new | amend | plate | return | cancel"""
    job.setdefault("events", []).append({"ts": _ts(), "type": kind, "detail": detail})
    if kind == "return":
        job["state"] = "returned"
    elif kind == "cancel":
        job["state"] = "cancelled"
    log_activity(kind, job, detail)


def _job_state(job):
    if job.get("state"):
        return job["state"]
    if job.get("cancelled"):
        return "cancelled"
    return "returned" if job.get("plate") == "📌" else "normal"


def _prune_jobs():
    """ตัดใบงานที่วันที่เก่าเกิน KEEP_DAYS วัน (ไม่ตัดใบที่ยังรอส่งในคิว) และจำกัดจำนวนไม่เกิน MAX_HISTORY"""
    global latest_jobs
    cutoff = _date_sort_key((now_th() - datetime.timedelta(days=KEEP_DAYS)).strftime("%d/%m/%Y"))
    queued = {id(j) for j in job_queue}
    latest_jobs = [j for j in latest_jobs
                   if id(j) in queued or _date_sort_key(_norm_job_date(j.get("date"))) >= cutoff][:MAX_HISTORY]



# ====================== ตาราง jobs ใน Supabase (1 แถว = 1 ใบงาน เก็บถาวร ค้นย้อนหลังได้) ======================
_sync_lock = threading.Lock()
_jobs_synced = {}       # uid -> ข้อมูลล่าสุดที่เขียนลงตารางแล้ว (เขียนซ้ำเฉพาะที่เปลี่ยน)
_jobs_warn_last = ""
_history_dates_cache = {"at": 0, "dates": []}


def _iso_date(d):
    try:
        return datetime.datetime.strptime(d, "%d/%m/%Y").strftime("%Y-%m-%d")
    except Exception:
        return None


def _job_row(job):
    d = _norm_job_date(job.get("date"))
    return {
        "uid": f"{d}|{_job_key(job)}",
        "job_date": _iso_date(d), "job_code": str(job.get("id", "")), "job_time": str(job.get("time", "")),
        "order_no": str(job.get("order", "")), "pickup": job.get("pickup", ""), "dropoff": job.get("dropoff", ""),
        "flight": str(job.get("flight", "")), "car_code": str(job.get("car_code", "")), "price": str(job.get("price", "")),
        "summary": job.get("formatted_summary", ""), "raw_text": job.get("text") or job.get("raw_text") or "",
        "status": job.get("status", ""), "state": _job_state(job), "plate": job.get("plate") or "",
        "events": job.get("events") or [], "received_at": job.get("received_at", ""),
    }


def sync_jobs_table(rows):
    """เขียนใบงานที่เปลี่ยนไปลงตาราง jobs (ถ้าตารางยังไม่ได้สร้าง จะแจ้งเตือนใน log แต่ระบบยังทำงานต่อได้)"""
    global _jobs_warn_last
    if not storage.using_supabase() or not rows:
        return
    with _sync_lock:
        changed = {}
        for r in rows:  # ใบงานซ้ำ uid เดียวกัน เอาอันแรก (ใหม่สุด)
            uid = r["uid"]
            if uid in changed:
                continue
            blob = json.dumps(r, sort_keys=True, ensure_ascii=False)
            if _jobs_synced.get(uid) != blob:
                changed[uid] = (r, blob)
        if not changed:
            return
        try:
            storage.upsert_rows("jobs", [v[0] for v in changed.values()], "uid")
            for uid, (_, blob) in changed.items():
                _jobs_synced[uid] = blob
            _jobs_warn_last = ""
        except storage.StorageError as e:
            if str(e) != _jobs_warn_last:
                _jobs_warn_last = str(e)
                print(f"⚠️ บันทึกตาราง jobs ไม่สำเร็จ (สร้างตารางด้วย supabase_setup.sql แล้วหรือยัง?): {e}")


def _rows_to_jobs(rows):
    jobs = []
    for r in rows:
        iso = r.get("job_date") or ""
        date = f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}" if len(iso) >= 10 else "-"
        jobs.append({
            "id": r.get("job_code", ""), "date": date, "time": r.get("job_time", ""), "order": r.get("order_no", ""),
            "pickup": r.get("pickup", ""), "pickup_display": r.get("pickup", ""),
            "dropoff": r.get("dropoff", ""), "dropoff_display": r.get("dropoff", ""),
            "flight": r.get("flight", ""), "car_code": r.get("car_code", ""), "price": r.get("price", ""),
            "formatted_summary": r.get("summary", ""), "text": r.get("raw_text", ""),
            "status": r.get("status", ""), "state": r.get("state") or "normal", "plate": r.get("plate", ""),
            "events": r.get("events") or [], "received_at": r.get("received_at", ""), "from_history": True,
            "uid": r.get("uid", ""),
        })
    return jobs


def history_dates():
    """วันที่ที่มีใบงานในตาราง jobs (เก็บผลไว้ 60 วินาที)"""
    if not storage.using_supabase():
        return []
    if time.time() - _history_dates_cache["at"] < 60:
        return _history_dates_cache["dates"]
    try:
        rows = storage.select_rows("jobs", {"select": "job_date", "order": "job_date.desc", "limit": "3000"})
    except Exception:  # ตาราง jobs ยังไม่พร้อม / storage.py เก่า: ข้ามไป ใช้เฉพาะข้อมูลในระบบ
        return _history_dates_cache["dates"]
    seen = []
    for r in rows:
        iso = r.get("job_date")
        if iso and iso not in seen:
            seen.append(iso)
    dates = [f"{i[8:10]}/{i[5:7]}/{i[0:4]}" for i in seen][:90]
    _history_dates_cache.update(at=time.time(), dates=dates)
    return dates


def persist_state():
    """บันทึกคิว + ประวัติใบงานลง Supabase (กันข้อมูลหายตอนเซิร์ฟเวอร์หลับ)"""
    with state_lock:
        snapshot = {
            "job_queue": list(job_queue),
            "latest_jobs": list(latest_jobs),
            "last_sender_id": last_sender_id,
            "timer_start_time": timer_start_time,
            "pending_update_dates": list(pending_update_dates),
        }
        job_rows = [_job_row(j) for j in latest_jobs]
    try:
        storage.set("job_state", snapshot)
    except storage.StorageError as e:
        print(f"⚠️ บันทึกคิวใบงานไม่สำเร็จ: {e}")

    global _logs_dirty
    with state_lock:
        dirty, activity, sent = _logs_dirty, list(_activity_log), list(_sent_log)
        _logs_dirty = False
    if dirty:
        try:
            storage.set("activity_log", activity)
            storage.set("sent_log", sent)
        except storage.StorageError as e:
            _logs_dirty = True
            print(f"⚠️ บันทึกประวัติไม่สำเร็จ: {e}")
    try:
        sync_jobs_table(job_rows)
    except Exception as e:  # ตาราง jobs เป็นของเสริม ห้ามทำให้ระบบหลักล้ม (เช่น อัป storage.py ไม่ครบ)
        print(f"⚠️ บันทึกตาราง jobs ไม่สำเร็จ: {e} (เช็คว่าอัป storage.py ล่าสุดขึ้น Render แล้ว)")


def restore_state():
    """โหลดคิวที่ค้างอยู่กลับมา และเริ่มนับเวลาต่อ"""
    global job_queue, latest_jobs, last_sender_id, timer_start_time, pending_update_dates
    try:
        snap = storage.get("job_state")
    except storage.StorageError as e:
        print(f"⚠️ โหลดคิวใบงานไม่สำเร็จ: {e}")
        return
    try:
        with state_lock:
            _activity_log[:] = storage.get("activity_log") or []
            _sent_log[:] = storage.get("sent_log") or []
    except storage.StorageError as e:
        print(f"⚠️ โหลดประวัติไม่สำเร็จ: {e}")
    if not snap:
        return
    with state_lock:
        job_queue = snap.get("job_queue") or []
        latest_jobs = snap.get("latest_jobs") or []
        last_sender_id = snap.get("last_sender_id")
        pending_update_dates = snap.get("pending_update_dates") or []
        if job_queue or pending_update_dates:
            wait_seconds = int(load_settings().get("wait_seconds", 10))
            started = snap.get("timer_start_time") or time.time()
            remaining = max(5, wait_seconds - (time.time() - started))
            _start_timer(remaining, started)
            print(f"♻️ กู้คืนคิวค้าง {len(job_queue)} ใบ จะส่งในอีก {int(remaining)} วินาที")


def _start_timer(seconds, started_at=None):
    global timer_thread, timer_start_time
    timer_start_time = started_at or time.time()
    timer_thread = threading.Timer(seconds, process_batch_jobs)
    timer_thread.daemon = True
    timer_thread.start()


# ====================== แปลงชื่อสถานที่ ======================
def _contains(text_lower, word):
    """เช็คคำ: ภาษาอังกฤษต้องเป็นคำเต็ม (กัน 'Chan' ไปตรงกับ 'Chang Klan'), ภาษาไทยเช็คแบบมีอยู่ในข้อความ"""
    w = str(word).strip().lower()
    if not w:
        return False
    if re.search(r"[a-z0-9]", w):
        return re.search(r"(?<![a-z0-9])" + re.escape(w) + r"(?![a-z0-9])", text_lower) is not None
    return w in text_lower


def _dict_candidates(d):
    """แปลงพจนานุกรม (dict หรือ list) เป็นรายการ (คำที่ใช้ค้นหา, ผลลัพธ์ภาษาไทย)"""
    out = []
    if isinstance(d, dict):
        for eng, th in d.items():
            eng_items = eng if isinstance(eng, list) else [eng]
            th_items = th if isinstance(th, list) else [th]
            th_items = [t for t in th_items if isinstance(t, str) and t.strip()]
            eng_items = [e for e in eng_items if isinstance(e, str) and e.strip()]
            result = th_items[0].strip() if th_items else (eng_items[0].strip() if eng_items else "")
            if not result:
                continue
            for word in eng_items + th_items:
                out.append((word, result))
    elif isinstance(d, list):
        for item in d:
            if not isinstance(item, dict):
                continue
            key = str(item.get("key", "")).strip()
            aliases = item.get("aliases", [])
            alias_list = aliases if isinstance(aliases, list) else [aliases]
            if not key:
                continue
            for word in [key] + [a for a in alias_list if isinstance(a, str)]:
                out.append((word, key))
    # คำยาวก่อน เพื่อให้ "New Petchaburi" ชนะ "Petchaburi"
    out.sort(key=lambda x: len(str(x[0])), reverse=True)
    return out


def _match_dict(text_lower, d):
    for word, result in _dict_candidates(d):
        if _contains(text_lower, word):
            return result
    return ""


def parse_location_rule_based(raw_location):
    """แปลงจุดรับ-จุดส่งด้วยกฎ และดึงพจนานุกรมล่าสุดจาก Settings"""
    if not raw_location or raw_location == "-":
        return "-"

    cleaned = raw_location.strip()
    upper_loc = cleaned.upper()
    lower_loc = cleaned.lower()

    settings = load_settings()
    current_main_roads = settings.get("main_roads_dict", DEFAULT_MAIN_ROADS_DICT)
    current_major_areas = settings.get("major_areas_dict", DEFAULT_MAJOR_AREAS_DICT)
    custom_keywords = settings.get("custom_keywords", [])

    # 0. Custom Keywords ที่ผู้ใช้ตั้งค่าเพิ่มเอง
    for item in custom_keywords:
        kw = str(item.get("keyword", "")).strip()
        target = str(item.get("zone", "")).strip()
        if kw and kw.lower() in lower_loc:
            return target if target else cleaned

    for name, target in BUILTIN_PLACES.items():
        if name in lower_loc:
            return target

    # 1. สนามบินดอนเมือง
    air_loc = re.sub(r"(?:KHET|KHWAENG|DISTRICT)\s+DON\s+MUEANG|(?:เขต|แขวง)ดอนเมือง", "", upper_loc)
    if any(k in air_loc for k in ["DMK", "DON MUEANG", "ดอนเมือง", "แอร์ดอน"]):
        return "แอร์ดอน"

    # 2. สนามบินสุวรรณภูมิ
    if (any(k in upper_loc for k in ["BKK", "SUVARNABHUMI", "SVB", "แอร์สุ", "สุวรรณภูมิ"])
            or re.search(r'(?<![A-Z])SUV(?![A-Z])', upper_loc)):
        return "แอร์สุ"

    # 3. Sukhumvit ตามด้วยเลขซอย -> สุขุมวิท22
    soi = re.search(r'\bSukhumvit\s*(?:Soi\s*)?(\d+)\b', cleaned, flags=re.IGNORECASE)
    if soi:
        return f"สุขุมวิท{soi.group(1)}"

    # 3.05 ซอยเพชรบุรี + เลขซอย -> เพชรบุรี 11
    pb = (re.search(r'\bSoi\s*(?:New\s*)?(?:Phetchaburi|Petchaburi)\s*(\d+)\b', cleaned, flags=re.IGNORECASE)
          or re.search(r'\b(?:Phetchaburi|Petchaburi)\s*Soi\s*(\d+)\b', cleaned, flags=re.IGNORECASE))
    if pb:
        return f"เพชรบุรี {pb.group(1)}"

    # 3.06 ชื่อไทยที่มี พระราม + เลข -> พระราม 5 / พระรามที่ 9 (คงรูปแบบตามที่เขียน)
    rm = re.search(r'พระราม\s*(ที่)?\s*(\d+)', cleaned)
    if rm:
        return f"พระราม{'ที่' if rm.group(1) else ''} {rm.group(2)}"

    # 3.1 เศษข้อความหลุด เช่น "ei Nuea" จาก Khlong Toei Nuea ให้ตีเป็นสุขุมวิท
    if "EI NUEA" in upper_loc or "KHLONG TOEI" in upper_loc:
        return "สุขุมวิท"

    # 4. ถนนหลัก แล้วค่อย 5. ย่านสำคัญ
    matched = _match_dict(lower_loc, current_main_roads) or _match_dict(lower_loc, current_major_areas)
    if matched:
        return matched

    # ไม่ตรงอะไรเลย ตัด "ซอย [ตัวเลข]" ออก
    cleaned_no_soi = re.sub(r'(?:ซอย|soi)\s*\d+', '', cleaned, flags=re.IGNORECASE).strip()
    return cleaned_no_soi if cleaned_no_soi else cleaned


def parse_job_line(line_text):
    # ป้องกันเคสที่ชื่อสถานที่/โรงแรมมีเครื่องหมาย - อยู่ข้างใน (เช่น Manhattan Hotel Bangkok-สุขุมวิท หรืออื่นๆ)
    # ให้มองหาเครื่องหมาย - ตัวสุดท้ายของบรรทัดเส้นทาง หรือเช็คเครื่องหมาย - ที่คั่นระหว่างจุดรับ-จุดส่งหลัก
    parts = line_text.split('-')
    
    # ถ้ามีมากกว่า 2 ส่วน (แสดงว่ามีเครื่องหมาย - เกินมาในชื่อสถานที่ เช่น Manhattan Hotel Bangkok-สุขุมวิท)
    if len(parts) > 2:
        # สมมติว่ารูปแบบคือ [จุดรับ] - [จุดส่ง] แต่จุดรับดันมีขีดคั่น ให้รวมตัวแรกกับตัวกลางเข้าด้วยกันเป็นจุดรับ
        pickup_raw = "-".join(parts[:-1]).strip()
        dropoff_raw = parts[-1].strip()
    elif len(parts) == 2:
        pickup_raw = parts[0].strip()
        dropoff_raw = parts[1].strip()
    else:
        return line_text, "-"
        
    pickup_clean = parse_location_rule_based(pickup_raw)
    dropoff_clean = parse_location_rule_based(dropoff_raw)
    
    return pickup_clean, dropoff_clean

def _place_code(mapped):
    """แปลงจุดรับ/จุดส่งเป็นรหัสย่อ: แอร์สุ=S, แอร์ดอน=D, อื่นๆ=BK"""
    if not mapped or mapped == "-":
        return "-"
    if mapped == "แอร์สุ":
        return "S"
    if mapped == "แอร์ดอน":
        return "D"
    return "BK"


_QUOTES = "\"'“”"
_SIGN_RE = re.compile(r'^ชูป้าย')
_DATE_LINE_RE = re.compile(r'^(\d{4})-(\d{1,2})-(\d{1,2})$|^(\d{1,2})/(\d{1,2})/(\d{4})$')
_TIME_LINE_RE = re.compile(r'^(?:Amend\s*>*\s*)?(\d{1,2})[:.](\d{2})$', re.IGNORECASE)
_FLIGHT_LINE_RE = re.compile(r'^[A-Za-z0-9]{2}\s?\d{1,4}[A-Za-z]?$')
_JOB_CODE_RE = re.compile(r'^[A-Za-z]{1,4}\d{1,4}$')
_ORDER_LINE_RE = re.compile(r'^(?=.*\d)(?=.*[A-Za-z])[A-Za-z0-9_-]{6,}$')
_PHONE_LINE_RE = re.compile(r'^\+?[\d\s()-]+$')
# บรรทัดทะเบียนรถ: CAR NO / Car No. / Car Plate / ทะเบียนรถ / ทะเบียน (ต้นบรรทัด มี : หรือไม่มีก็ได้)
_PLATE_RE = re.compile(r'^[ \t]*(?:CAR\s*NO\.?|CAR\s*PLATE|ทะเบียนรถ|ทะเบียน)[ \t]*[:：\-]?[ \t]*(\S.*)$',
                       re.IGNORECASE | re.MULTILINE)
_DRIVER_LINE_RE = re.compile(
    r'^\s*(?:(?:CAR\s*NO|CAR\s*PLATE|ทะเบียนรถ|ทะเบียน)\.?\s*[:：\-]?|(?:TEL|APP|TYPE\s*CAR)\.?\s*[:：])', re.IGNORECASE)


def _clean_lines(raw_text):
    return [ln.strip().strip(_QUOTES).strip() for ln in raw_text.splitlines()]


def _norm_date(text):
    m = _DATE_LINE_RE.match(text.strip())
    if not m:
        return None
    if m.group(1):
        y, mo, d = m.group(1), m.group(2), m.group(3)
    else:
        d, mo, y = m.group(4), m.group(5), m.group(6)
    return f"{int(d):02d}/{int(mo):02d}/{y}"


def _norm_time(text):
    m = re.match(r'^(?:Amend\s*>*\s*)?(\d{1,2})[:.](\d{2})', str(text).strip(), re.IGNORECASE)
    return f"{int(m.group(1)):02d}:{m.group(2)}" if m else text


def is_voucher_format(raw_text):
    """ใบงานแบบใหม่ (KKday / Book on A Budget ...) มีหัวใบงาน: เที่ยวบิน / วันที่ / เวลา / ชูป้าย ... / รหัสงาน"""
    return "【" not in raw_text and any(_SIGN_RE.match(ln) for ln in _clean_lines(raw_text))


def extract_plate(raw_text):
    """ดึงป้ายทะเบียนจากบรรทัด CAR NO : 5 ขพ 9425 -> 5ขพ9425"""
    m = _PLATE_RE.search(raw_text or "")
    if not m:
        return ""
    return re.sub(r'\s+', '', m.group(1))


def strip_driver_lines(raw_text):
    """ตัดบรรทัดข้อมูลคนขับ (CAR NO / TEL / APP / TYPE CAR) ออก เหลือแต่ใบงาน"""
    return "\n".join(ln for ln in raw_text.splitlines() if not _DRIVER_LINE_RE.match(ln)).strip()


def _parse_voucher(raw_text):
    lines = [ln for ln in _clean_lines(raw_text) if ln]
    header = {}
    used = set()

    sign_idx = next(i for i, ln in enumerate(lines) if _SIGN_RE.match(ln))
    used.add(sign_idx)
    if sign_idx + 1 < len(lines) and _JOB_CODE_RE.match(lines[sign_idx + 1]):
        header["id"] = lines[sign_idx + 1].upper()
        used.add(sign_idx + 1)
    # ย้อนขึ้นไปจากบรรทัดชูป้าย: เวลา -> วันที่ -> เที่ยวบิน
    i = sign_idx - 1
    if i >= 0 and _TIME_LINE_RE.match(lines[i]):
        header["time"] = _norm_time(lines[i])
        used.add(i)
        i -= 1
    if i >= 0 and _norm_date(lines[i]):
        header["date"] = _norm_date(lines[i])
        used.add(i)
        i -= 1
    if i >= 0 and _FLIGHT_LINE_RE.match(lines[i]):
        header["flight"] = re.sub(r'\s+', '', lines[i]).upper()
        used.add(i)

    ticket = [ln for i, ln in enumerate(lines) if i not in used]

    order = next((ln for ln in ticket if _ORDER_LINE_RE.match(ln)), "-")

    pickup_raw, dropoff_raw = "-", "-"
    phone_idx = next((i for i, ln in enumerate(ticket)
                      if _PHONE_LINE_RE.match(ln) and len(re.sub(r'\D', '', ln)) >= 8), None)
    if phone_idx is not None and phone_idx + 1 < len(ticket):
        # หลังเบอร์โทร: บรรทัดแรก = จุดรับ ที่เหลือ = จุดส่ง (ที่อยู่อาจยาวหลายบรรทัด)
        pickup_raw = ticket[phone_idx + 1]
        dropoff_raw = " ".join(ticket[phone_idx + 2:]) or "-"
    elif len(ticket) >= 2:
        pickup_raw, dropoff_raw = ticket[-2], ticket[-1]

    car_info = match_car("\n".join(ticket))

    return {
        "id": header.get("id", ""),
        "date": header.get("date", "-"),
        "time": header.get("time", "-"),
        "flight": header.get("flight", "-"),
        "pickup_raw": pickup_raw,
        "dropoff_raw": dropoff_raw,
        "car_info": car_info,
        "order": order,
    }


_KV_ORDER_RE = re.compile(r'^\s*(?:Booking Reference ID|Booking Reference Number|Booking ID)\s*[:：]', re.IGNORECASE | re.MULTILINE)


def is_kv_format(raw_text):
    """ใบงานแบบ 'คีย์ : ค่า' (Booking Reference ID / Booking ID ...) เช่น TT1, Million Holiday"""
    return bool(_KV_ORDER_RE.search(raw_text or "")) and "【" not in raw_text


def _kv_pairs(raw_text):
    pairs = {}
    for ln in raw_text.splitlines():
        m = re.match(r'^\s*([A-Za-z][A-Za-z /]*?)\s*[:：]\s*(.*)$', ln)
        if m:
            key = re.sub(r'\s+', ' ', m.group(1).strip().lower())
            pairs.setdefault(key, m.group(2).strip())
    return pairs


def _kv_get(pairs, *keys):
    for k in keys:
        if pairs.get(k):
            return pairs[k]
    return ""


def _kv_place(name, address):
    """เลือกข้อความสถานที่ที่จะแปลง: ชื่อสถานที่ก่อน (ตัดที่อยู่ออก) ถ้าไม่ตรงอะไรค่อยรวมที่อยู่ช่วย"""
    name = (name or "").strip()
    address = (address or "").strip()
    if not name and not address:
        return "-"
    short = re.split(r'\s\d', name)[0].strip() or name
    if parse_location_rule_based(short) != short:
        return short
    return f"{name} {address}".strip()


def _parse_kv(raw_text):
    lines = [ln.strip() for ln in raw_text.splitlines() if ln.strip()]
    pairs = _kv_pairs(raw_text)

    job_id = ""
    if lines and _JOB_CODE_RE.match(lines[0]):
        job_id = lines[0].upper()
    elif lines and _JOB_CODE_RE.match(lines[-1]):
        job_id = lines[-1].upper()

    date_raw = _kv_get(pairs, "date request", "date of use", "date", "pick up date")
    date_val = _norm_date(date_raw) or "-"
    if date_val == "-" and date_raw:
        for fmt in ("%d %b %Y", "%d %B %Y"):
            try:
                date_val = datetime.datetime.strptime(date_raw, fmt).strftime("%d/%m/%Y")
                break
            except ValueError:
                pass

    time_raw = _kv_get(pairs, "pick up time", "pickup time", "time request", "time")
    time_val = "-"
    tm = re.search(r'(\d{1,2})[:.](\d{2})', time_raw)
    if tm:
        hour = int(tm.group(1))
        if re.search(r'PM', time_raw, re.IGNORECASE) and hour < 12:
            hour += 12
        time_val = f"{hour:02d}:{tm.group(2)}"

    flight_raw = _kv_get(pairs, "flight number", "flight no", "flight")
    if not flight_raw:  # บรรทัดแบบ "Flight no.  EK385" (ไม่มี :)
        fm = re.search(r'^\s*flight\s*(?:no\.?|number)?\s*[:：]?\s*([A-Za-z0-9][A-Za-z0-9 \-]*)$',
                       raw_text, re.IGNORECASE | re.MULTILINE)
        flight_raw = fm.group(1) if fm else ""
    flight = re.sub(r'[^A-Za-z0-9]', '', flight_raw).upper() or "-"

    pickup_raw = _kv_place(_kv_get(pairs, "pick up location", "pickup location"),
                           _kv_get(pairs, "pick up address", "pickup address"))
    dropoff_raw = _kv_place(_kv_get(pairs, "drop off location", "dropoff location"),
                            _kv_get(pairs, "drop off address", "dropoff address"))
    if pairs.get("hotel name"):  # แบบ Airport to Hotel / Hotel to Airport
        airport = _kv_get(pairs, "airport of arrival", "airport of departure", "airport")
        hotel = _kv_place(pairs["hotel name"], _kv_get(pairs, "address"))
        if re.search(r'hotel\s*to\s*airport', raw_text, re.IGNORECASE):
            pickup_raw, dropoff_raw = hotel, airport or "-"
        else:
            pickup_raw, dropoff_raw = airport or "-", hotel

    car_info = None
    for key in ("participant", "package", "car type", "vehicle"):
        car_info = match_car(pairs.get(key, ""))
        if car_info:
            break

    return {
        "id": job_id,
        "date": date_val,
        "time": time_val,
        "flight": flight,
        "pickup_raw": pickup_raw,
        "dropoff_raw": dropoff_raw,
        "car_info": car_info,
        "order": _kv_get(pairs, "booking reference id", "booking reference number", "booking id") or "-",
    }


_WEEKDAY_DATE_RE = re.compile(r'^(?:Amend\s*>*\s*)?(?:[A-Za-z]+,\s*)?(\d{1,2})\s+([A-Za-z]{3,9})\s+(\d{4})$', re.IGNORECASE)
_ROUTE_LINE_RE = re.compile(r'^[A-Za-z][A-Za-z ]*\s-\s[A-Za-z][A-Za-z ]*$')  # เช่น BKK - Don Mueang / Pak Kret District - DMK
_FLIGHT_SPACED_RE = re.compile(r'^([A-Za-z0-9]{2})\s*-\s*(\d{1,4})$')


def is_plain_format(raw_text):
    """ใบงานแบบไม่มีหัวข้อ: บรรทัดสุดท้ายเป็นรหัสงาน และมีบรรทัดวันที่ เช่น 'Saturday, 17 Oct 2026'"""
    if "【" in raw_text:
        return False
    lines = [ln.strip() for ln in raw_text.splitlines() if ln.strip()]
    return (len(lines) >= 8 and bool(_JOB_CODE_RE.match(lines[-1]))
            and any(_WEEKDAY_DATE_RE.match(ln) for ln in lines))


def _short_name(text):
    """ชื่อสถานที่แบบสั้น: ตัดวงเล็บ ที่อยู่ และตัวเลขท้ายออก"""
    text = text.split("(")[0]
    return re.split(r',|\s\d', text)[0].strip() or text.strip()


def _parse_plain(raw_text):
    lines = [ln.strip().strip(_QUOTES).strip() for ln in raw_text.splitlines()]
    lines = [ln for ln in lines if ln]
    job_id = lines[-1].upper()

    date_idx = next(i for i, ln in enumerate(lines) if _WEEKDAY_DATE_RE.match(ln))
    dm = _WEEKDAY_DATE_RE.match(lines[date_idx])
    date_val = "-"
    for fmt in ("%d %b %Y", "%d %B %Y"):
        try:
            date_val = datetime.datetime.strptime(f"{dm.group(1)} {dm.group(2)} {dm.group(3)}", fmt).strftime("%d/%m/%Y")
            break
        except ValueError:
            pass

    time_val = "-"
    if date_idx + 1 < len(lines) and _TIME_LINE_RE.match(lines[date_idx + 1]):
        time_val = _norm_time(lines[date_idx + 1])

    flight, flight_idx = "-", date_idx
    if date_idx >= 1:
        fm = _FLIGHT_SPACED_RE.match(lines[date_idx - 1]) or _FLIGHT_LINE_RE.match(lines[date_idx - 1])
        if fm:
            flight = re.sub(r'[\s-]+', '', lines[date_idx - 1]).upper()
            flight_idx = date_idx - 1

    # หลังบรรทัดเส้นทาง (BKK - Don Mueang) ถึงบรรทัดเที่ยวบิน/วันที่ เรียงเป็น
    # [จุดรับ, หมายเหตุจุดรับ, จุดส่ง, หมายเหตุจุดส่ง, ...] โดย "-" = ไม่มีหมายเหตุ (หมายเหตุเช่น "terminal 2 ต่างประเทศ" ไม่ใช่สถานที่)
    route_idx = next((i for i, ln in enumerate(lines[:flight_idx]) if _ROUTE_LINE_RE.match(ln)), None)
    block = lines[route_idx + 1:flight_idx] if route_idx is not None else []
    if len(block) >= 3:
        places = [block[0], block[2]]
    else:  # รูปแบบสั้น: เอาเฉพาะบรรทัดที่ไม่ใช่ "-"
        places = [ln for ln in block if ln != "-"]
    def _plain_place(text):
        short = _short_name(text)
        if parse_location_rule_based(short) != short.strip():
            return short
        if parse_location_rule_based(text) != text.strip():
            return text
        return short

    pickup_raw = _plain_place(places[0]) if places and places[0] != "-" else "-"
    dropoff_raw = _plain_place(places[1]) if len(places) > 1 and places[1] != "-" else "-"

    car_info = match_car(lines[date_idx + 2]) if date_idx + 2 < len(lines) else None
    if not car_info:
        car_info = match_car("\n".join(lines[date_idx + 1:-1]))

    order = next((ln for ln in lines[:1] if _ORDER_LINE_RE.match(ln)), "-")
    return {
        "id": job_id, "date": date_val, "time": time_val, "flight": flight,
        "pickup_raw": pickup_raw, "dropoff_raw": dropoff_raw,
        "car_info": car_info, "order": order,
    }


def _parse_legacy(raw_text):
    # รหัสใบงาน = หัวข้อบรรทัดแรกของใบงาน (เช่น F03) รหัสอื่นก็ใช้ได้หมด
    job_id = ""
    for ln in raw_text.splitlines():
        ln = ln.strip()
        if not ln:
            continue
        # บรรทัดแรกต้องเป็นรหัสสั้นๆ ไม่ใช่บรรทัดข้อมูลที่ขึ้นต้นด้วย 【
        if "【" not in ln and "】" not in ln and len(ln) <= 20:
            job_id = ln
        break
    if not job_id:
        id_match = re.search(r'\b(F\d+)\b', raw_text, re.IGNORECASE)
        job_id = id_match.group(1).strip() if id_match else ""

    date_match = re.search(r'(?:【(?:日期วันที่|日期|วันที่)】|วันที่|Date)[:\s]*([\d/\-]+)', raw_text, re.IGNORECASE)
    date_val = date_match.group(1).strip() if date_match else "-"

    time_match = re.search(r'(?:【(?:เวลาเวลา|เวลา|เวลา)】|เวลา|Time)[:\s]*([\d:]+)', raw_text, re.IGNORECASE)
    if not time_match:
        time_match = re.search(r'(\d{2}:\d{2})', raw_text)
    time_val = time_match.group(1).strip() if time_match else "-"

    flight_val = "-"
    flight_match = re.search(r'(?:【(?:航班flight|航班|flight)】|เที่ยวบิน|flight|Flight)[:\s]*([A-Za-z0-9]+)', raw_text, re.IGNORECASE)
    if not flight_match:
        flight_match = re.search(r'✈️?\s*([A-Za-z]{2}\d+|\d{3,4})', raw_text)
    if flight_match:
        flight_val = flight_match.group(1).strip()

    # ดึงค่าจากแท็ก 【接รับ】 และ 【ส่ง】 โดยตรง 100% ตามรูปแบบในรูปภาพ
    pickup_raw = "-"
    dropoff_raw = "-"

    pickup_tag = re.search(r'【接รับ】\s*(.+)', raw_text)
    if pickup_tag:
        pickup_raw = pickup_tag.group(1).strip()

    dropoff_tag = re.search(r'【送ส่ง】\s*(.+)', raw_text)
    if dropoff_tag:
        dropoff_raw = dropoff_tag.group(1).strip()

    # จัดการเรื่องขนาดรถ: ดูบรรทัดขนาดรถก่อน ถ้าไม่มีค่อยหาทั้งใบงาน
    car_raw_match = re.search(r'(?:【(?:车型ขนาดรถ|车型|ขนาดรถ)】|รถ|ขนาดรถ|Car)[:\s]*(.+)', raw_text, re.IGNORECASE)
    car_info = match_car(car_raw_match.group(1)) if car_raw_match else None
    if not car_info:
        car_info = match_car(raw_text)

    order_match = re.search(r'(?:【(?:客户订单号|订单号|Order)】|Order|Order Number|คำสั่งซื้อ|Order ID)[:\s]*([0-9A-Za-z_-]+)', raw_text, re.IGNORECASE)
    if not order_match:
        order_match = re.search(r'\b(\d{8,20})\b', raw_text)
    order_val = order_match.group(1).strip() if order_match else "-"

    return {
        "id": job_id,
        "date": date_val,
        "time": time_val,
        "flight": flight_val,
        "pickup_raw": pickup_raw,
        "dropoff_raw": dropoff_raw,
        "car_info": car_info,
        "order": order_val,
    }


def build_summary(job_id, mark, time_val, price, car_code, route_code, pickup, dropoff, flight, order):
    return f"{job_id}({mark}){time_val}/{price}#{car_code}/{route_code}\n{pickup}-{dropoff} ✈️{flight}\n{order}"


def apply_plate(summary, plate):
    """แทนที่ไอคอนในวงเล็บแรกด้วยป้ายทะเบียน: F01(🥶)... -> F01(5ขพ9425)..."""
    return re.sub(r'^([^(\n]*)\([^)\n]*\)', lambda m: f"{m.group(1)}({plate})", summary, count=1)


def parse_job_text(raw_text, fallback_id="F01"):
    if not raw_text:
        return {}

    if is_kv_format(raw_text):
        fields = _parse_kv(raw_text)
    elif is_plain_format(raw_text):
        fields = _parse_plain(raw_text)
    elif is_voucher_format(raw_text):
        fields = _parse_voucher(raw_text)
    else:
        fields = _parse_legacy(raw_text)
    job_id = fields["id"] or fallback_id
    time_val = _norm_time(fields["time"])
    flight_val = fields["flight"]
    order_val = fields["order"]
    pickup_raw = fields["pickup_raw"]
    dropoff_raw = fields["dropoff_raw"]

    # แปลงชื่อสถานที่
    pickup_mapped = parse_location_rule_based(pickup_raw)
    dropoff_mapped = parse_location_rule_based(dropoff_raw)

    car_info = fields["car_info"] or {"code": "5S", "price": "380"}
    car_code = car_info["code"]
    price_val = car_info["price"]

    # เงื่อนไขไอคอน: ถ้าตำแหน่งที่ 2 (จุดส่ง) เป็น แอร์ดอน หรือ แอร์สุ ให้ใช้ 🔥 ที่เหลือใช้ 🥶
    if dropoff_mapped in ["แอร์ดอน", "แอร์สุ"]:
        icon_symbol = "🔥"
    else:
        icon_symbol = "🥶"

    # รหัสจุดรับ-จุดส่ง: S = แอร์สุ, D = แอร์ดอน, BK = สถานที่/โรงแรมในกรุงเทพ
    route_code = f"{_place_code(pickup_mapped)}-{_place_code(dropoff_mapped)}"

    formatted_summary = build_summary(job_id, icon_symbol, time_val, price_val, car_code, route_code,
                                      pickup_mapped, dropoff_mapped, flight_val, order_val)

    return {
        "id": job_id,
        "date": fields["date"],
        "time": time_val,
        "pickup_raw": pickup_raw,
        "pickup": pickup_mapped,
        "dropoff_raw": dropoff_raw,
        "dropoff": dropoff_mapped,
        "flight": flight_val,
        "car_code": car_code,
        "price": price_val,
        "order": order_val,
        "formatted_summary": formatted_summary
    }


def _job_key(job):
    order = str(job.get("order", "")).strip()
    if order and order != "-":
        return "order:" + order
    return "text:" + str(job.get("text", "")).strip()


def _find_by_order(pool, order, date=None):
    """ใบงานที่เลข Order ตรงกัน ถ้ารู้วันที่ให้เลือกของวันเดียวกันก่อน
    (Order เดียวกันต่างวัน = คนละงาน เช่น ขาไปและขากลับในการจองเดียว) ถ้าไม่มีวันเดียวกันค่อยใช้ Order อย่างเดียว"""
    order = str(order or "").strip()
    if not order or order == "-":
        return []
    same = [j for j in pool if str(j.get("order", "")).strip() == order]
    if date and date != "-":
        d = _norm_job_date(date)
        exact = [j for j in same if _norm_job_date(j.get("date")) == d]
        if exact:
            return exact
    return same


def _job_uid(job):
    """รหัสเฉพาะของใบงาน (วันที่|คีย์) ใช้อ้างอิงตอนลบและเป็นคีย์ในตาราง jobs"""
    return f"{_norm_job_date(job.get('date'))}|{_job_key(job)}"


def _order_ok(job):
    o = str(job.get("order", "")).strip()
    return bool(o) and o != "-"


def _functional_sig(job):
    return tuple(str(job.get(k, "")).strip().lower()
                 for k in ("id", "time", "pickup", "dropoff", "flight", "car_code", "price")) + (_norm_job_date(job.get("date")),)


def _same_job(a, b):
    """ใบงานเดียวกันไหม: Order เดียวกัน หรือ (ทั้งคู่ไม่มี Order) รหัส/วันที่/เวลา/จุดรับ-ส่ง/เที่ยวบิน/รถ/ราคาตรงกันหมด"""
    if _job_key(a) == _job_key(b):
        return True
    if not _order_ok(a) and not _order_ok(b):
        return _functional_sig(a) == _functional_sig(b)
    return False


def add_job_to_queue(text, sender_id=None):
    """เพิ่มใบงานเข้าคิว คืน (สำเร็จไหม, ข้อมูลใบงาน หรือ ข้อความ)"""
    global latest_jobs, last_sender_id

    with state_lock:
        # รหัสใบงานตามลำดับในคิวปัจจุบัน (F01, F02, ...) กัน F ชนกัน
        used_ids = {j.get("id") for j in job_queue}
        seq = len(job_queue) + 1
        while f"F{seq:02d}" in used_ids:
            seq += 1
        parsed_info = parse_job_text(text, fallback_id=f"F{seq:02d}")

        now = now_th()
        job_item = {
            "id": parsed_info["id"],
            "date": parsed_info["date"] if parsed_info["date"] != "-" else now.strftime("%d/%m/%Y"),
            "text": text,
            "raw_text": text,
            "time": parsed_info["time"] if parsed_info["time"] != "-" else now.strftime("%H:%M:%S"),
            "pickup": parsed_info["pickup"],
            "pickup_display": parsed_info["pickup"],
            "dropoff": parsed_info["dropoff"],
            "dropoff_display": parsed_info["dropoff"],
            "flight": parsed_info["flight"],
            "order": parsed_info["order"],
            "car_code": parsed_info["car_code"],
            "price": parsed_info["price"],
            "formatted_summary": parsed_info["formatted_summary"],
            "received_at": now.strftime("%Y-%m-%d %H:%M:%S"),
            "status": "รอส่ง"
        }

        # กันใบงานซ้ำ (เช่น LINE ส่ง webhook ซ้ำ หรือกดส่งซ้ำ)
        key = _job_key(job_item)
        new_date = _norm_job_date(job_item["date"])
        is_amend = bool(re.search(r'amend', text, re.IGNORECASE))
        if any(_same_job(j, job_item) and _norm_job_date(j.get("date")) == new_date for j in job_queue):
            log_activity("duplicate", job_item, "พบใบงานซ้ำในคิว (ไม่รับ)")
            return False, "ℹ️ ใบงานนี้อยู่ในคิวแล้ว (Order เดียวกัน) รอบอทสรุปสักครู่"

        # เคยรับใบงานนี้แล้ว (Order เดียวกัน): ส่งซ้ำข้อความเดิมเป๊ะ = ไม่รับ ถ้าแก้ไขแล้ว = รับ แต่คงป้ายทะเบียน/📌 เดิมไว้
        prev = next((j for j in latest_jobs
                     if _same_job(j, job_item) and j.get("status") != "ยกเลิก"
                     and (_norm_job_date(j.get("date")) == new_date or is_amend or j.get("raw_text") == text)), None)
        if prev:
            same_fields = all(str(prev.get(k, "")) == str(job_item.get(k, ""))
                              for k in ("id", "date", "time", "pickup", "dropoff", "flight", "car_code", "price"))
            if prev.get("raw_text") == text and prev.get("status") == "ส่งแล้ว" and same_fields:
                log_activity("duplicate", prev, "ส่งข้อความเดิมซ้ำ (ไม่รับ)")
                return False, "ℹ️ ใบงานนี้รับไว้แล้ว (ข้อความเดิมซ้ำ) สรุปล่าสุด:\n\n" + str(prev.get("formatted_summary", ""))
            if prev.get("plate"):
                job_item["plate"] = prev["plate"]
                job_item["formatted_summary"] = apply_plate(job_item["formatted_summary"], prev["plate"])
            if prev.get("price_override") not in (None, ""):
                job_item["price_override"] = prev["price_override"]
                job_item["formatted_summary"] = apply_price_value(job_item["formatted_summary"], prev["price_override"])
            if prev.get("price_adj"):
                job_item["price_adj"] = prev["price_adj"]
                job_item["formatted_summary"] = apply_price_adj(job_item["formatted_summary"], prev["price_adj"])
            if prev.get("edits"):
                keep = {}
                for f, v in prev["edits"].items():
                    if str(job_item.get(f)) == str((prev.get("edit_orig") or {}).get(f)):  # ใบงานใหม่ไม่ได้เปลี่ยนค่านี้ = คงที่เราแก้ไว้
                        keep[f] = v
                if keep:
                    job_item.update(keep)
                    job_item["edits"] = keep
                    job_item["edit_orig"] = {f: prev["edit_orig"][f] for f in keep if f in prev.get("edit_orig", {})}
                    job_item["formatted_summary"] = rebuild_summary(job_item)
            job_item["state"] = _job_state(prev)
            if job_item["state"] == "cancelled":
                job_item["cancelled"] = True
                job_item["formatted_summary"] = f"{job_item['id']} ❌CANCEL❌"
            job_item["events"] = list(prev.get("events", []))
            latest_jobs = [j for j in latest_jobs if j is not prev]  # ใบงานแก้ไข: แทนที่ใบเดิม ไม่ให้ซ้ำในรายการ
        job_item.setdefault("state", "normal")

        job_queue.append(job_item)
        if sender_id:
            last_sender_id = sender_id

        latest_jobs.insert(0, job_item)
        mark_job(job_item, "amend" if prev else "new", "ใบงานแก้ไข" if prev else "รับใบงานใหม่")
        _prune_jobs()

        _ensure_timer()

    persist_state()
    return True, job_item


def _ensure_timer():
    """เริ่มนับเวลาส่งรอบใหม่ (ถ้ายังไม่ได้นับอยู่) ต้องเรียกตอนถือ state_lock"""
    if timer_thread is None:
        wait_seconds = int(load_settings().get("wait_seconds", 10))
        _start_timer(wait_seconds)
        print(f"⏱️ เริ่มนับเวลาประมวลผลอีก {wait_seconds} วินาที...")


def _norm_job_date(raw_date):
    raw_date = str(raw_date or now_th().strftime("%d/%m/%Y"))
    try:
        return datetime.datetime.strptime(raw_date.replace('-', '/'), "%d/%m/%Y").strftime("%d/%m/%Y")
    except Exception:
        return raw_date


def update_plate(text, plate, sender_id=None):
    """อัปเดตใบงานเดิมเป็นป้ายทะเบียน (ค้นจาก Order) ถ้าไม่เจอใบงานเดิมจะเพิ่มเป็นใบงานใหม่"""
    global last_sender_id
    ticket = strip_driver_lines(text)
    parsed = parse_job_text(ticket)
    order = str(parsed.get("order", "")).strip()

    with state_lock:
        matches = _find_by_order(job_queue + latest_jobs, order, parsed.get("date"))
        if matches:
            queued_keys = {id(j) for j in job_queue}
            only_queued = True
            for job in matches:
                job["plate"] = plate
                job["formatted_summary"] = apply_plate(job["formatted_summary"], plate)
                mark_job(job, "return" if plate == RETURN_MARK else "plate", f"{plate} (CAR NO + ใบงาน จาก …{str(sender_id or '-')[-6:]})")
                if id(job) not in queued_keys and job.get("status") != "รอส่ง":
                    only_queued = False
            # ใบงานที่ส่งไปแล้ว -> ส่งใบสรุปของวันนั้นใหม่พร้อมป้ายทะเบียน
            if not only_queued:
                d = _norm_job_date(matches[0].get("date"))
                if d not in pending_update_dates:
                    pending_update_dates.append(d)
                if sender_id:
                    last_sender_id = sender_id
                _ensure_timer()
            job_id = matches[0].get("id")
        else:
            job_id = None

    if job_id is None:
        ok, result = add_job_to_queue(ticket, sender_id=sender_id)
        if not ok:
            return False, result
        with state_lock:
            result["plate"] = plate
            result["formatted_summary"] = apply_plate(result["formatted_summary"], plate)
            mark_job(result, "return" if plate == RETURN_MARK else "plate", f"{plate} (CAR NO + ใบงาน ไม่พบใบงานเดิม จาก …{str(sender_id or '-')[-6:]})")
        persist_state()
        return True, {"id": result["id"], "plate": plate, "new": True,
                      "formatted_summary": result.get("formatted_summary", "")}

    persist_state()
    return True, {"id": job_id, "plate": plate}


_DATE_TOKEN = r'(\d{1,2})(?:[/\-.](\d{1,2}))?(?:[/\-.](\d{4}))?'
# รูปแบบ: "F22 คืน", "F22 คืน 16", "F22 คืน 16/10", "F22 16/10 คืน"
_RETURN_RE = re.compile(
    r'^\s*([A-Za-z]+\d+)\s*(?:' + _DATE_TOKEN + r'\s*คืน|คืน(?:\s*' + _DATE_TOKEN + r')?)\s*$')
RETURN_MARK = "📌"
_RETURN_TICKET_RE = re.compile(r'^\s*📌\s*คืน[ \t]*\r?\n(.+)$', re.DOTALL)


def parse_return_command(text):
    """คืน (รหัสใบงาน, วัน, เดือน, ปี) หรือ None ถ้าไม่ใช่คำสั่งคืนป้าย (วัน/เดือน/ปี ว่างได้)"""
    m = _RETURN_RE.match(text or "")
    if not m:
        return None
    g = m.groups()
    day, mon, year = next(((g[i], g[i + 1], g[i + 2]) for i in (1, 4) if g[i]), (None, None, None))
    return g[0], day, mon, year


def return_plate(job_id, day=None, month=None, year=None, sender_id=None):
    """'VK06 คืน' -> เปลี่ยนป้ายของ VK06 เป็น 📌 แล้วส่งรายการย่อของวันนั้นใหม่
    ถ้ารหัสเดียวกันมีหลายวัน ต้องระบุวันที่ (เช่น 'F22 คืน 16' หรือ 'F22 16/10 คืน') ไม่งั้นจะไม่แก้อะไร"""
    global last_sender_id
    job_id = job_id.strip().upper()
    with state_lock:
        matches = [j for j in job_queue + latest_jobs
                   if str(j.get("id", "")).strip().upper() == job_id and j.get("status") != "ยกเลิก"]
        if day:
            def _same_day(j):
                parts = _norm_job_date(j.get("date")).split("/")
                if len(parts) != 3:
                    return False
                return (int(parts[0]) == int(day)
                        and (not month or int(parts[1]) == int(month))
                        and (not year or int(parts[2]) == int(year)))
            matches = [j for j in matches if _same_day(j)]
        if not matches:
            return False, f"ไม่พบใบงาน {job_id}" + (f" ของวันที่ {day}" + (f"/{month}" if month else "") if day else "")
        dates = sorted({_norm_job_date(j.get("date")) for j in matches}, key=_date_sort_key)
        if len(dates) > 1:
            return False, (f"{job_id} มีหลายวัน ({', '.join(dates)}) กรุณาระบุวัน เช่น "
                           f"{job_id} คืน {dates[0][:2]}")
        queued_keys = {id(j) for j in job_queue}
        already_sent = False
        for job in matches:
            job["plate"] = RETURN_MARK
            job["formatted_summary"] = apply_plate(job["formatted_summary"], RETURN_MARK)
            mark_job(job, "return", RETURN_MARK)
            if id(job) not in queued_keys and job.get("status") != "รอส่ง":
                already_sent = True
        if already_sent:
            if dates[0] not in pending_update_dates:
                pending_update_dates.append(dates[0])
            if sender_id:
                last_sender_id = sender_id
            _ensure_timer()
    persist_state()
    return True, {"id": job_id, "date": dates[0]}


_CANCEL_RE = re.compile(r'^\W*(?:cancel(?:led|lation)?|ยกเลิก)\W*$', re.IGNORECASE | re.MULTILINE)


def is_cancel_ticket(text):
    """ใบงานที่มีบรรทัด 'cancel' (หรือ 'ยกเลิก') ต่อท้าย = แจ้งยกเลิกงานนั้น"""
    return bool(_CANCEL_RE.search(text or ""))


def cancel_job(text, sender_id=None):
    """ใบงาน + บรรทัด cancel -> สรุปย่อของงานนั้นเป็น 'F09 ❌CANCEL❌' (ค้นจากเลข Order ถ้าไม่มีใช้รหัส+วันที่)"""
    global last_sender_id
    ticket = _CANCEL_RE.sub("", text).strip()
    parsed = parse_job_text(ticket)
    order = str(parsed.get("order", "")).strip()
    mark = lambda jid: f"{jid} ❌CANCEL❌"

    with state_lock:
        pool = job_queue + latest_jobs
        if order and order != "-":
            matches = _find_by_order(pool, order, parsed.get("date"))
        else:
            d = _norm_job_date(parsed.get("date"))
            matches = [j for j in pool if str(j.get("id", "")).strip().upper() == str(parsed.get("id", "")).strip().upper()
                       and _norm_job_date(j.get("date")) == d]
        if matches:
            queued_keys = {id(j) for j in job_queue}
            already_sent = False
            for job in matches:
                job["cancelled"] = True
                job["formatted_summary"] = mark(job.get("id"))
                mark_job(job, "cancel", "ยกเลิกงาน")
                if id(job) not in queued_keys and job.get("status") != "รอส่ง":
                    already_sent = True
            job_id = matches[0].get("id")
            if already_sent:
                d = _norm_job_date(matches[0].get("date"))
                if d not in pending_update_dates:
                    pending_update_dates.append(d)
                if sender_id:
                    last_sender_id = sender_id
                _ensure_timer()
        else:
            job_id = None

    if job_id is None:  # ไม่เจองานเดิม: เพิ่มเป็นรายการ CANCEL ใหม่
        ok, result = add_job_to_queue(ticket, sender_id=sender_id)
        if not ok:
            return False, result
        with state_lock:
            result["cancelled"] = True
            result["formatted_summary"] = mark(result["id"])
            mark_job(result, "cancel", "ยกเลิกงาน (ไม่พบใบงานเดิม)")
        persist_state()
        return True, result

    persist_state()
    return True, {"id": job_id, "cancelled": True}


_UPDATE_RE = re.compile(r'^\s*(?:update|อั[พป]?เด[ทต])\s*(.*?)\s*$', re.IGNORECASE)


def parse_update_date(text):
    """วันที่ในคำสั่ง Update: '' / วันนี้ = วันนี้, พรุ่งนี้, 17, 17/10, 17/10/2026 (ไม่ระบุเดือน/ปี ใช้ของวันนี้) คืน dd/mm/yyyy หรือ None"""
    text = (text or "").strip().lower()
    now = now_th()
    if text in ("", "วันนี้", "today"):
        return now.strftime("%d/%m/%Y")
    if text in ("พรุ่งนี้", "tomorrow"):
        return (now + datetime.timedelta(days=1)).strftime("%d/%m/%Y")
    m = re.match(r'^(\d{1,2})(?:[/\-.](\d{1,2}))?(?:[/\-.](\d{2,4}))?$', text)
    if not m:
        return None
    day = int(m.group(1))
    month = int(m.group(2)) if m.group(2) else now.month
    year = int(m.group(3)) if m.group(3) else now.year
    if year < 100:
        year += 2000
    if year > 2400:  # พ.ศ.
        year -= 543
    try:
        return datetime.date(year, month, day).strftime("%d/%m/%Y")
    except ValueError:
        return None


def jobs_for_date(date):
    """ใบงานของวันนั้น (ไม่ซ้ำ) จากระบบ ถ้าไม่มีลองอ่านจากตาราง jobs (ใบงานเก่า)"""
    seen, out = set(), []
    with state_lock:
        for j in latest_jobs:
            if j.get("status") == "ยกเลิก" or _norm_job_date(j.get("date")) != date:
                continue
            k = _job_key(j)
            if k not in seen:
                seen.add(k)
                out.append(j)
    if not out and storage.using_supabase() and _iso_date(date):
        try:
            rows = storage.select_rows("jobs", {"job_date": f"eq.{_iso_date(date)}", "select": "*", "limit": "1000"})
            out = _rows_to_jobs(rows)
        except Exception as e:
            print(f"⚠️ อ่านตาราง jobs ไม่สำเร็จ: {e}")
    return out


# ====================== คำสั่งในแชท: ปิดใบงาน / แก้ไขราคา / ตัดตาราง / NOW ======================
def _reply_long(conf, reply_token, text):
    """ตอบข้อความยาว: แบ่งเป็นหลายข้อความ (ไม่เกิน 5) โดยตัดที่บรรทัดว่าง ใช้ reply จึงไม่เสียโควต้า"""
    chunks, cur = [], ""
    for block in text.split("\n\n"):
        piece = block if not cur else cur + "\n\n" + block
        if len(piece) > 4900 and cur:
            chunks.append(cur)
            cur = block
        else:
            cur = piece
    if cur:
        chunks.append(cur)
    chunks = [c[:4900] for c in chunks][:5]
    try:
        with ApiClient(conf) as api_client:
            MessagingApi(api_client).reply_message(ReplyMessageRequest(
                reply_token=reply_token, messages=[TextMessage(text=c) for c in chunks]))
    except Exception as e:
        print(f"⚠️ ตอบกลับไม่ได้: {e}")


def _job_candidates(job_id, date_text=None):
    """ใบงานทุกใบของรหัสนั้น (+วันที่ถ้าระบุ) ในวันที่เลือก (ใหม่สุดก่อน) คืน (รายการ, ข้อความผิดพลาด)
    รหัสซ้ำหลายวันและไม่ระบุวัน: ใช้ของวันนี้ถ้ามี ไม่งั้นให้ระบุวัน"""
    job_id = job_id.strip().upper()
    want = None
    if date_text:
        want = parse_update_date(date_text)
        if want is None:
            return [], f"อ่านวันที่ '{date_text}' ไม่ได้"
    with state_lock:
        seen, cands = set(), []
        for j in latest_jobs:
            if j.get("status") == "ยกเลิก" or str(j.get("id", "")).strip().upper() != job_id:
                continue
            d = _norm_job_date(j.get("date"))
            if want and d != want:
                continue
            k = (d, _job_key(j))
            if k in seen:
                continue
            seen.add(k)
            cands.append(j)
    if not cands:
        return [], f"ไม่พบใบงาน {job_id}" + (f" ของวันที่ {want}" if want else "")
    dates = sorted({_norm_job_date(j.get("date")) for j in cands}, key=_date_sort_key)
    if len(dates) > 1:
        today = now_th().strftime("%d/%m/%Y")
        if today not in dates:
            return [], f"{job_id} มีหลายวัน ({', '.join(dates)}) กรุณาระบุวัน เช่น {job_id} {dates[0][:2]}"
        dates = [today]
    return [j for j in cands if _norm_job_date(j.get("date")) == dates[0]], None


def resolve_job(job_id, date_text=None):
    """หาใบงานจากรหัส (+วันที่ถ้าระบุ) คืน (job, ข้อความผิดพลาด) ถ้ามีหลายใบรหัสเดียวกันวันเดียวกัน ใช้ใบใหม่สุด"""
    cands, err = _job_candidates(job_id, date_text)
    if err:
        return None, err
    return cands[0], None


def _queue_compact_update(job, sender_id=None):
    """ใบงานที่ส่งไปแล้วถูกแก้ -> ส่งรายการย่อของวันนั้นใหม่ (เรียกตอนถือ state_lock)"""
    global last_sender_id
    if any(j is job for j in job_queue) or job.get("status") == "รอส่ง":
        return
    d = _norm_job_date(job.get("date"))
    if d not in pending_update_dates:
        pending_update_dates.append(d)
    if sender_id:
        last_sender_id = sender_id
    _ensure_timer()


def apply_price_adj(summary, adj):
    """ใส่ส่วนปรับราคาต่อท้ายราคา: /380#... -> /380+50#..."""
    first, sep, rest = summary.partition("\n")
    adj_str = f"{adj:+d}" if adj else ""
    first = re.sub(r'/(\d+)(?:[+-]\d+)?#', lambda m: f"/{m.group(1)}{adj_str}#", first, count=1)
    return first + sep + rest


def apply_price_value(summary, value):
    """ตั้งราคาในใบสรุปเป็นค่าใหม่ (ลบส่วนปรับเดิมออก): /380+50#... -> /200#..."""
    first, sep, rest = summary.partition("\n")
    first = re.sub(r'/(\d+)(?:[+-]\d+)?#', lambda m: f"/{value}#", first, count=1)
    return first + sep + rest


def _effective_price(job):
    """ราคาที่ใช้จริง: ราคาที่ตั้งใหม่ (price_override) ถ้ามี ไม่งั้นราคาที่อ่านจากใบงาน"""
    return job.get("price_override") if job.get("price_override") not in (None, "") else job.get("price", "")


def set_price_override(job, value, sender_id=None):
    job["price_override"] = value
    job["price_adj"] = 0
    job["formatted_summary"] = apply_price_value(job["formatted_summary"], value)
    mark_job(job, "price", f"ตั้งราคา {value}")
    _queue_compact_update(job, sender_id)


def set_price_adj(job, adj, sender_id=None):
    job["price_adj"] = adj
    job["formatted_summary"] = apply_price_adj(job["formatted_summary"], adj)
    mark_job(job, "price", f"{adj:+d}" if adj else "ยกเลิกส่วนปรับราคา")
    _queue_compact_update(job, sender_id)


_EMOJI_MARKS = ("🥶", "🔥", "📌")


def has_plate(job):
    """มีทะเบียนแล้วไหม: ในวงเล็บแรกของสรุปมีข้อความที่ไม่ใช่ 🥶 / 🔥 / 📌 (รองรับ 9027, กต9027, กต 9027)"""
    m = re.match(r'^[^(\n]*\(([^)\n]*)\)', str(job.get("formatted_summary", "")))
    content = m.group(1).strip() if m else ""
    return bool(content) and content not in _EMOJI_MARKS


def set_plate_manual(job, plate, detail, sender_id=None):
    job["plate"] = plate
    job["formatted_summary"] = apply_plate(job["formatted_summary"], plate)
    mark_job(job, "return" if plate == RETURN_MARK else "plate", f"{plate} ({detail})")
    _queue_compact_update(job, sender_id)


_CLOSE_RE = re.compile(r'^\s*(?:ปิด)?ใบงาน\s+([A-Za-z]+\d+)(?:\s+(\S+))?\s*$')


_DRIVER_KEY_RE = re.compile(
    r'^\s*(?:(?:CAR\s*NO\.?|CAR\s*PLATE|ทะเบียนรถ|ทะเบียน)\s*[:：\-]?\s*\S|(?:NAME|TEL|APP|TYPE\s*CAR)\s*[:：])', re.IGNORECASE)


def split_driver_block(text):
    """แยกข้อมูลคนขับ (NAME / CAR NO / TEL / APP / TYPE CAR) ออกจากใบงาน โดยยึดบรรทัด CAR NO เป็นหลัก
    ข้อความอื่นที่อยู่ก่อนบล็อกคนขับ (เช่น คำว่า Update) ไม่สนใจ
    คืน (บล็อกคนขับ, ใบงาน, ตำแหน่งที่ส่งมา 'top'/'bottom'/None) ไม่เจอ CAR NO คืน ("", ข้อความเดิม, None)"""
    lines = text.strip().splitlines()
    a = next((i for i, ln in enumerate(lines) if _PLATE_RE.match(ln)), None)
    if a is None:
        return "", text.strip(), None
    while a > 0 and _DRIVER_KEY_RE.match(lines[a - 1]):  # NAME : ... ที่อยู่เหนือ CAR NO
        a -= 1
    b = a
    while b < len(lines) and (not lines[b].strip() or _DRIVER_KEY_RE.match(lines[b])):
        b += 1
    block = "\n".join(lines[a:b]).strip()
    before, after = "\n".join(lines[:a]).strip(), "\n".join(lines[b:]).strip()
    if _is_ticket_text(after):
        return block, after, "top"
    if _is_ticket_text(before):
        return block, before, "bottom"
    return block, (after or before), None


def merge_close_text(job_code, driver_block, ticket):
    """รวมข้อความปิดใบงาน: รหัส F... ข้อมูลคนขับ (ทะเบียน) ไว้บน ใบงานแบบอื่นไว้ล่าง"""
    if re.match(r'^F\d', str(job_code or "").strip().upper()):
        return driver_block.strip() + "\n\n" + ticket.strip()
    return ticket.strip() + "\n\n" + driver_block.strip()


def _is_ticket_text(text):
    return ("【客户订单号】" in text or is_voucher_format(text) or is_kv_format(text) or is_plain_format(text))


def close_with_ticket(text, sender_id=None):
    """ปิดใบงานด้วยข้อความเดียว: ใบงาน + ข้อมูลคนขับ -> ใส่ทะเบียนตามเลข Order
    คืน (สำเร็จไหม, ข้อความตอบ) ข้อความตอบว่าง = ไม่ต้องตอบกลับ (ที่ส่งมาจัดลำดับถูกอยู่แล้ว)
    ตอบข้อความรวมเฉพาะเมื่อต้องจัดลำดับใหม่: รหัส F ข้อมูลคนขับไว้บน รหัสอื่นไว้ล่าง"""
    block, ticket, position = split_driver_block(text)
    plate = extract_plate(block)
    if not block or not plate:
        return False, "ไม่พบข้อมูลคนขับ (บรรทัด CAR NO : ...) ในข้อความ"
    if not _is_ticket_text(ticket):
        return False, "ไม่พบใบงานในข้อความ แนบใบงานทั้งใบ หรือพิมพ์ ใบงาน F01 ไว้บรรทัดแรก"
    ok, result = update_plate(ticket, plate, sender_id=sender_id)
    if not ok:
        return False, str(result)
    wanted = "top" if re.match(r'^F\d', str(result.get("id") or "").strip().upper()) else "bottom"
    if result.get("new"):
        return True, merge_close_text(result.get("id"), block, ticket) + "\n\n➕ ไม่พบใบงานเดิม เพิ่มเป็นใบงานใหม่แล้ว"
    if position == wanted:
        return True, ""
    return True, merge_close_text(result.get("id"), block, ticket)


def close_job(job_id, date_text, block, sender_id=None):
    """ปิดใบงาน: ข้อมูลคนขับ (มี CAR NO) -> ใส่ทะเบียนในใบงาน และคืน 'ข้อมูลคนขับ + ใบงานต้นฉบับ'"""
    plate = extract_plate(block)
    if not plate:
        return False, "ไม่พบบรรทัด CAR NO : ... ในข้อความ"
    job, err = resolve_job(job_id, date_text)
    if err:
        return False, err
    with state_lock:
        set_plate_manual(job, plate, "ปิดใบงาน", sender_id)
        original = job.get("text") or job.get("raw_text") or job.get("formatted_summary", "")
    persist_state()
    return True, merge_close_text(job.get("id"), block, str(original))


_PRICE_HEAD_RE = re.compile(r'^\s*(?:แก้ไข|อั[พป]?เด[ทต])ราคา\s*$')
_PRICE_LINE_RE = re.compile(
    r'^\s*([A-Za-z]+\d+)\s*(?:(\d{1,2}(?:[/\-.]\d{1,2}(?:[/\-.]\d{2,4})?)?)\s+)?([+-]?)\s*(\d+)\s*$')


def price_command(lines, sender_id=None):
    """แก้ไขราคา: แต่ละบรรทัด 'F01 +50' (หรือ 'F01 17/10 +50') ตั้งส่วนปรับราคาของใบงานนั้น เก็บถาวรในใบงาน"""
    results = []
    for ln in lines:
        if not ln.strip():
            continue
        m = _PRICE_LINE_RE.match(ln)
        if not m:
            results.append(f"⚠️ อ่านไม่ออก: {ln.strip()}")
            continue
        job_id, date_text, sign, num = m.groups()
        if not sign and not date_text and int(num) < 100:  # "F01 19" น่าจะเป็นวันที่ที่ลืมใส่ราคา ไม่ใช่ราคา 19
            results.append(f"⚠️ อ่านไม่ออก: {ln.strip()} (ถ้า {num} คือวันที่ ให้ใส่ราคาตามหลัง เช่น {job_id} {num} 200 หรือ {job_id} {num} +50)")
            continue
        job, err = resolve_job(job_id, date_text)
        if err:
            results.append(f"⚠️ {err}")
            continue
        with state_lock:
            if sign:   # +50 / -30 : ปรับต่อท้ายราคาเดิม
                set_price_adj(job, int(num) * (1 if sign == "+" else -1), sender_id)
            else:      # 200 : ใช้ราคาที่แก้ไข
                set_price_override(job, int(num), sender_id)
            first = str(job["formatted_summary"]).splitlines()[0]
        results.append(f"✅ {first}")
    persist_state()
    return "\n".join(results) or "ไม่พบรายการราคา เช่น\nแก้ไขราคา\nF01 +50  (เพิ่มจากราคาเดิม)\nF02 200  (ใช้ราคา 200)"


_UNCANCEL_RE = re.compile(r'^\s*(?:uncancel|un)\s+([A-Za-z]+\d+)(?:\s+(\S+))?\s*$', re.IGNORECASE)
_UNCANCEL_LINE_RE = re.compile(r'^\W*(?:uncancel|un)\W*$', re.IGNORECASE | re.MULTILINE)


def _restore_cancelled(job, sender_id=None):
    """ดึงงานที่ยกเลิกกลับมา: สร้างใบสรุปจากข้อความต้นฉบับใหม่ แล้วใส่ทะเบียน/ส่วนปรับราคาเดิมกลับ"""
    with state_lock:
        raw = job.get("text") or job.get("raw_text") or ""
        summary = parse_job_text(raw, fallback_id=str(job.get("id", "")))["formatted_summary"] if raw else ""
        if not summary:
            return None
        plate = job.get("plate")
        if job.get("edits"):
            summary = rebuild_summary(job)
        else:
            if plate:
                summary = apply_plate(summary, plate)
            if job.get("price_override") not in (None, ""):
                summary = apply_price_value(summary, job["price_override"])
            if job.get("price_adj"):
                summary = apply_price_adj(summary, job["price_adj"])
        job["formatted_summary"] = summary
        job["cancelled"] = False
        job["state"] = "returned" if plate == RETURN_MARK else "normal"
        mark_job(job, "uncancel", "ดึงงานกลับมา")
        _queue_compact_update(job, sender_id)
    persist_state()
    return summary


def _uncancel_reply(job, summary, remaining=0):
    date = _norm_job_date(job.get("date"))
    text_date = _norm_job_date(parse_job_text(job.get("text") or job.get("raw_text") or "").get("date"))
    msg = f"✅ ดึงงานกลับมาแล้ว (วันที่ {date})\n\n{summary}"
    if str(summary).split("\n")[0].endswith("---"):
        msg += "\n\n⚠️ ข้อมูลใบงานนี้ไม่ครบ (จุดรับ/จุดส่งว่าง เพราะใบที่ยกเลิกเป็นข้อความสั้น) ส่งใบงานเต็มเข้ามาใหม่เพื่อแก้"
    if text_date and text_date != date and text_date != "-":
        msg += f"\n\n⚠️ ข้อความต้นฉบับของใบงานนี้ระบุวันที่ {text_date} แต่ระบบบันทึกเป็นวันที่ {date} ตรวจสอบด้วย"
    if remaining:
        msg += f"\n\nℹ️ ยังมี {job.get('id')} ที่ยกเลิกอยู่อีก {remaining} ใบ พิมพ์ uncancel {job.get('id')} อีกครั้งเพื่อดึงใบถัดไป"
    return msg


def _do_uncancel(cancelled, others_active_note, sender_id):
    job = cancelled[0]
    summary = _restore_cancelled(job, sender_id)
    if not summary:
        return False, f"{job.get('id')} ไม่มีข้อความต้นฉบับให้สร้างใบสรุปใหม่"
    return True, _uncancel_reply(job, summary, remaining=len(cancelled) - 1)


def _is_stub_cancel(job):
    """รายการยกเลิกที่สร้างจากข้อความสั้น (ไม่มีจุดรับ/จุดส่ง) สร้างใบสรุปกลับมาไม่ได้"""
    raw = job.get("text") or job.get("raw_text") or ""
    return not raw or str(parse_job_text(raw, fallback_id=str(job.get("id", "")))["formatted_summary"]).split("\n")[0].endswith("---")


def uncancel_job(job_id, date_text, sender_id=None):
    """uncancel F09 / uncancel F09 18 -> ดึงงานที่ยกเลิกกลับมา
    ถ้ามีหลายใบรหัสเดียวกันวันเดียวกัน (ปกติ + ยกเลิก) เลือกเฉพาะใบที่ยกเลิกอยู่ (ใหม่สุดก่อน)
    รายการยกเลิกที่เป็นข้อความสั้น (กู้ใบสรุปไม่ได้) และมีใบปกติรหัส/วันเดียวกันอยู่แล้ว จะถูกลบทิ้งแทนการกู้"""
    cands, err = _job_candidates(job_id, date_text)
    if err:
        return False, err
    cancelled = [j for j in cands if _job_state(j) == "cancelled"]
    active = [j for j in cands if _job_state(j) != "cancelled"]
    if not cancelled:
        return False, f"{cands[0].get('id')} ไม่ได้ถูกยกเลิก"
    note = ""
    stubs = [j for j in cancelled if _is_stub_cancel(j)]
    if stubs and active:
        with state_lock:
            for j in stubs:
                j["status"] = "ยกเลิก"  # ลบออกจากรายการ
                mark_job(j, "uncancel", "ลบรายการยกเลิกที่เป็นข้อความสั้น (มีใบปกติอยู่แล้ว)")
            _queue_compact_update(active[0], sender_id)
        persist_state()
        cancelled = [j for j in cancelled if j not in stubs]
        note = (f"🧹 ลบรายการ {active[0].get('id')} ❌CANCEL❌ ที่เป็นข้อความสั้นแล้ว {len(stubs)} รายการ "
                f"(วันนี้มี {active[0].get('id')} ปกติอยู่แล้ว {len(active)} ใบ)")
        if not cancelled:
            return True, note
    ok, msg = _do_uncancel(cancelled, None, sender_id)
    return ok, (note + "\n\n" + msg) if note and ok else msg


def uncancel_ticket(text, sender_id=None):
    """ใบงานเต็ม + บรรทัด uncancel (หรือ un) -> ทำให้ใบงานนี้กลับมาเป็นงานปกติ ใช้ข้อความใบงานที่แนบมาสร้างใบสรุปใหม่
    1) มีใบเดิม (Order+วันเดียวกัน) ที่ยกเลิกอยู่ -> ดึงกลับ  2) ไม่เจอตาม Order แต่มีรายการ CANCEL สั้น ๆ ของรหัส/วันเดียวกัน -> เปลี่ยนเป็นใบงานนี้
    3) ไม่มีเลย -> เพิ่มเป็นใบงานใหม่  รายการ CANCEL สั้น ๆ ที่เหลือของรหัส/วันเดียวกันถูกลบทิ้ง"""
    ticket = _UNCANCEL_LINE_RE.sub("", text).strip()
    parsed = parse_job_text(ticket)
    order = str(parsed.get("order", "")).strip()
    code = str(parsed.get("id", "")).strip().upper()
    date = _norm_job_date(parsed.get("date"))
    notes = []
    with state_lock:
        found = [j for j in _find_by_order(latest_jobs, order, parsed.get("date")) if j.get("status") != "ยกเลิก"]
        same_code = [j for j in latest_jobs if j.get("status") != "ยกเลิก"
                     and str(j.get("id", "")).strip().upper() == code and _norm_job_date(j.get("date")) == date]
        target = next((j for j in found if _job_state(j) == "cancelled"), None)
        if target is None and not found:
            # ไม่เจอตาม Order: ใช้รายการ CANCEL สั้น ๆ ของรหัส/วันเดียวกันที่ใหม่สุด
            target = next((j for j in same_code if _job_state(j) == "cancelled" and _is_stub_cancel(j)), None)
        active = [j for j in found if _job_state(j) != "cancelled"]
        if target is not None:
            target["text"] = target["raw_text"] = ticket       # ใช้ข้อความเต็มที่แนบมา
            target["order"] = order if order and order != "-" else target.get("order", "")
            for k in ("pickup", "dropoff", "flight", "car_code", "price", "time", "date"):
                if parsed.get(k) not in (None, "", "-"):
                    target[k] = parsed[k]
            target["pickup_display"], target["dropoff_display"] = target.get("pickup", ""), target.get("dropoff", "")
    summary = None
    if target is not None:
        summary = _restore_cancelled(target, sender_id)
    elif not found:
        ok, result = add_job_to_queue(ticket, sender_id=sender_id)
        if not ok:
            return False, str(result)
        notes.append("➕ ไม่พบใบงานเดิม เพิ่มเป็นใบงานใหม่แล้ว")
        target, summary = result, result.get("formatted_summary", "")
    else:
        notes.append(f"ℹ️ {code} ใบนี้เป็นงานปกติอยู่แล้ว ไม่ได้ถูกยกเลิก")
        target = found[0]
    # ลบรายการ CANCEL สั้น ๆ ที่เหลือของรหัส/วันเดียวกัน (ใบปกติของงานนี้มีแล้ว)
    with state_lock:
        leftovers = [j for j in same_code if j is not target and _job_state(j) == "cancelled" and _is_stub_cancel(j)]
        for j in leftovers:
            j["status"] = "ยกเลิก"
            mark_job(j, "uncancel", "ลบรายการยกเลิกที่เป็นข้อความสั้น (มีใบเต็มแล้ว)")
        if leftovers:
            _queue_compact_update(target, sender_id)
    if leftovers:
        persist_state()
        notes.append(f"🧹 ลบรายการ {code} ❌CANCEL❌ ที่เป็นข้อความสั้นแล้ว {len(leftovers)} รายการ")
    if summary:
        return True, "\n\n".join([_uncancel_reply(target, summary)] + notes) if not notes or notes[0][0] != "➕" else "\n\n".join(notes + [summary])
    return True, "\n\n".join(notes)


_EDIT_HEAD_RE = re.compile(r'^\s*แก้ไข(?:ใบงาน|ใบสรุป|สรุป)\s*$')
_EDIT_LINE_RE = re.compile(
    r'^\s*([A-Za-z]+\d+)\s*(?:(\d{1,2}(?:[/\-.]\d{1,2}(?:[/\-.]\d{2,4})?)?)\s+)?(\S.*?)\s*(?:>|＞|→|->)\s*(\S.*?)\s*$')
_TIME_ANY_RE = re.compile(r'^\d{1,2}[:.]\d{2}$')
_AIRPORT_ALIASES = {"แอร์สุ": "แอร์สุ", "สุวรรณภูมิ": "แอร์สุ", "bkk": "แอร์สุ",
                    "แอร์ดอน": "แอร์ดอน", "ดอนเมือง": "แอร์ดอน", "dmk": "แอร์ดอน"}


def is_edit_message(lines):
    """ทุกบรรทัด (ไม่นับหัว 'แก้ไขใบงาน') เป็นรูปแบบ  รหัส [วัน] ค่าเดิม>ค่าใหม่"""
    body = [ln for ln in lines if ln.strip()]
    if body and _EDIT_HEAD_RE.match(body[0]):
        body = body[1:]
    return bool(body) and all(_EDIT_LINE_RE.match(ln) for ln in body)


def rebuild_summary(job):
    """สร้างใบสรุปใหม่จากข้อมูลของใบงาน (เวลา/เส้นทาง/รถ/Order/ราคา/ทะเบียน) ใช้ตอนแก้ไขใบสรุป"""
    icon = "🔥" if job.get("dropoff") in ("แอร์ดอน", "แอร์สุ") else "🥶"
    mark = job.get("plate") or icon
    adj = job.get("price_adj") or 0
    price = f"{_effective_price(job)}{adj:+d}" if adj else f"{_effective_price(job)}"
    route = f"{_place_code(job.get('pickup', '-'))}-{_place_code(job.get('dropoff', '-'))}"
    return build_summary(str(job.get("id", "")), mark, job.get("time") or "-", price, job.get("car_code") or "5S", route,
                         job.get("pickup") or "-", job.get("dropoff") or "-", job.get("flight") or "-", job.get("order") or "-")


def _car_by_code(text):
    t = str(text).strip().lower()
    for v in CAR_PRICING_MAP.values():
        if v["code"].lower() == t:
            return v
    return None


def _norm_place(name):
    n = name.strip()
    if n == "แอร์":
        raise ValueError("ระบุให้ชัดว่า แอร์สุ หรือ แอร์ดอน")
    return _AIRPORT_ALIASES.get(n.lower(), n)


def _split_route(text):
    """'แอร์สุ-บางกะปิ' -> ('แอร์สุ', 'บางกะปิ') รองรับสนามบินที่ต้นหรือท้าย และชื่อสถานที่ที่มี - อยู่ข้างใน"""
    t = text.strip()
    for a in ("แอร์สุ", "แอร์ดอน"):
        if t.startswith(a + "-"):
            return a, t[len(a) + 1:].strip()
        if t.endswith("-" + a):
            return t[:-len(a) - 1].strip(), a
    if "-" in t:
        p, d = t.split("-", 1)
        return _norm_place(p), _norm_place(d)
    raise ValueError("เส้นทางต้องเป็นรูปแบบ จุดรับ-จุดส่ง")


def _apply_edit(job, old, new, sender_id=None):
    """แก้ 1 ค่าในใบงาน (เวลา / Order / ขนาดรถ / เส้นทาง) ตรวจว่า 'ค่าเดิม' ตรงกับที่ระบบมีก่อน คืน (สำเร็จไหม, ข้อความ)"""
    code = job.get("id")
    if _job_state(job) == "cancelled":
        return False, f"{code} ถูกยกเลิกอยู่ ใช้ uncancel ก่อน"
    edits, orig = job.setdefault("edits", {}), job.setdefault("edit_orig", {})
    cur_order = str(job.get("order", "")).strip()
    old_uid = None

    def remember(field):
        orig.setdefault(field, job.get(field))

    if _TIME_ANY_RE.match(old):
        if not _TIME_ANY_RE.match(new):
            return False, f"เวลาใหม่ '{new}' รูปแบบไม่ถูก (เช่น 20:00)"
        if _norm_time(old) != _norm_time(job.get("time", "")):
            return False, f"เวลาเดิมของ {code} คือ {job.get('time')} ไม่ตรงกับ {old}"
        remember("time")
        job["time"] = edits["time"] = _norm_time(new)
        kind = "เวลา"
    elif cur_order not in ("", "-") and old == cur_order:
        if re.search(r'\s', new):
            return False, "Order ใหม่ต้องไม่มีช่องว่าง"
        old_uid = _job_uid(job)
        remember("order")
        job["order"] = edits["order"] = new
        kind = "Order"
    elif _car_by_code(old) or _car_by_code(new):
        car = _car_by_code(new)
        if not car:
            return False, f"ไม่รู้จักขนาดรถ '{new}'"
        if str(job.get("car_code", "")).lower() != old.lower():
            return False, f"ขนาดรถเดิมของ {code} คือ {job.get('car_code')} ไม่ตรงกับ {old}"
        remember("car_code"), remember("price")
        job["car_code"], job["price"] = car["code"], car["price"]
        edits["car_code"], edits["price"] = car["code"], car["price"]
        kind = "ขนาดรถ"
    elif "-" in old or "-" in new:
        try:
            old_p, old_d = _split_route(old)
            new_p, new_d = _split_route(new)
        except ValueError as e:
            return False, str(e)
        if (old_p, old_d) != (job.get("pickup"), job.get("dropoff")):
            return False, f"เส้นทางเดิมของ {code} คือ {job.get('pickup')}-{job.get('dropoff')} ไม่ตรงกับ {old}"
        remember("pickup"), remember("dropoff")
        job["pickup"] = job["pickup_display"] = edits["pickup"] = new_p
        job["dropoff"] = job["dropoff_display"] = edits["dropoff"] = new_d
        kind = "เส้นทาง"
    else:
        return False, (f"ไม่พบค่าเดิม '{old}' ใน {code} (เวลา {job.get('time')} • รถ {job.get('car_code')} • "
                       f"เส้นทาง {job.get('pickup')}-{job.get('dropoff')} • Order {cur_order or '-'})")
    job["formatted_summary"] = rebuild_summary(job)
    mark_job(job, "edit", f"{kind}: {old} → {new}")
    _queue_compact_update(job, sender_id)
    if old_uid:  # Order เปลี่ยน = uid เปลี่ยน ลบแถวเก่าในตาราง jobs
        _jobs_synced.pop(old_uid, None)
        try:
            storage.delete_rows("jobs", "uid", old_uid)
        except Exception as e:
            print(f"⚠️ ลบแถวเก่าในตาราง jobs ไม่สำเร็จ: {e}")
    return True, ""


def edit_job_command(lines, sender_id=None):
    """แก้ไขใบสรุป: แต่ละบรรทัด 'VK01 19 13:25>20:00' / 'VK01 19 แอร์สุ-บางกะปิ>แอร์สุ-สีลม' / 'VK01 19 5S>EV-SUV' / 'VK01 19 ORDER>ORDERใหม่'"""
    body = [ln for ln in lines if ln.strip()]
    if body and _EDIT_HEAD_RE.match(body[0]):
        body = body[1:]
    results = []
    for ln in body:
        m = _EDIT_LINE_RE.match(ln)
        if not m:
            results.append(f"⚠️ อ่านไม่ออก: {ln.strip()}")
            continue
        job_id, date_text, old, new = m.groups()
        job, err = resolve_job(job_id, date_text)
        if err:
            results.append(f"⚠️ {err}")
            continue
        with state_lock:
            ok, msg = _apply_edit(job, old.strip(), new.strip(), sender_id)
            first = str(job.get("formatted_summary", "")).splitlines()[0] if ok else ""
        results.append(f"✅ {first}" if ok else f"⚠️ {msg}")
    persist_state()
    return "\n".join(results)


_CLEAR_PLATE_RE = re.compile(r'^\s*(?:ล้าง|ลบ)ทะเบียน\s+([A-Za-z]+\d+)(?:\s+(\S+))?\s*$')


def clear_plate(job_id, date_text, sender_id=None):
    """ล้างทะเบียน: เอาทะเบียน/📌 ออก กลับเป็นไอคอนเดิม (🔥 จุดส่งเป็นสนามบิน นอกนั้น 🥶)"""
    job, err = resolve_job(job_id, date_text)
    if err:
        return False, err
    with state_lock:
        icon = "🔥" if job.get("dropoff") in ("แอร์ดอน", "แอร์สุ") else "🥶"
        job["plate"] = ""
        job["formatted_summary"] = apply_plate(job["formatted_summary"], icon)
        if job.get("state") == "returned":
            job["state"] = "normal"
        mark_job(job, "plate", "ล้างทะเบียน")
        _queue_compact_update(job, sender_id)
        first = str(job["formatted_summary"]).splitlines()[0]
    persist_state()
    return True, f"✅ ล้างทะเบียนแล้ว: {first}"


_FULL_TABLE_RE = re.compile(r'^\s*ดูตารางเต็ม\s*(.*?)\s*$')
_VIEW_JOB_RE = re.compile(r'^\s*ดูงาน\s+([A-Za-z]+\d+)(?:\s*[-\s]\s*(\S+))?\s*$')


def original_summary(job):
    """ใบสรุปเต็มของงาน โดยวงเล็บเป็นไอคอนเดิมก่อนใส่ทะเบียน (🔥 จุดส่งเป็นสนามบิน นอกนั้น 🥶) งานที่ยกเลิกแสดง CANCEL"""
    summary = str(job.get("formatted_summary", "")).strip()
    if _job_state(job) == "cancelled":
        return summary
    icon = "🔥" if job.get("dropoff") in ("แอร์ดอน", "แอร์สุ") else "🥶"
    return apply_plate(summary, icon)


def full_table_text(date_text):
    """ดูตารางเต็ม: ใบสรุปเต็มทุกงานของวันนั้น เรียงตามรหัส วงเล็บเป็นไอคอนเดิม"""
    date = parse_update_date(date_text)
    if date is None:
        return "รูปแบบ: ดูตารางเต็ม หรือ ดูตารางเต็ม 18 หรือ ดูตารางเต็ม 18/10"
    jobs = sorted(jobs_for_date(date), key=_id_sort_key)
    if not jobs:
        return f"ไม่มีใบงานวันที่ {date}"
    return f"📅 {date}\n\n" + "\n\n".join(original_summary(j) for j in jobs)


def view_job_text(job_id, date_text):
    """ดูงาน F01-18: ใบสรุปเต็มของงานเดียว วงเล็บเป็นไอคอนเดิม"""
    job, err = resolve_job(job_id, date_text)
    if err:
        return f"⚠️ {err}"
    return original_summary(job)


_CUT_RE = re.compile(r'^\s*ตัดตาราง\s*(.*?)\s*$')
_NOW_RE = re.compile(r'^\s*(?:now|นาว)\s+([A-Za-z]+\d+)(?:\s+(\S+))?\s*$', re.IGNORECASE)


def cut_table_text(date_text):
    """ตัดตาราง: ใบสรุปเต็มเฉพาะงานที่ยังไม่มีทะเบียน (รวมงานที่คืน 📌) ไม่รวมงานที่ยกเลิก"""
    date = parse_update_date(date_text)
    if date is None:
        return "รูปแบบ: ตัดตาราง หรือ ตัดตาราง 18 หรือ ตัดตาราง 18/10"
    jobs = [j for j in sorted(jobs_for_date(date), key=_id_sort_key)
            if _job_state(j) != "cancelled" and not has_plate(j)]
    if not jobs:
        return f"📅 {date}\n\nทุกงานมีทะเบียนแล้ว ✅"
    return f"📅 {date}\n\n" + "\n\n".join(str(j["formatted_summary"]).strip() for j in jobs)


def now_text(job_id, date_text):
    """งานด่วน: 'now ขนาดรถ' / 'จุดรับ-จุดส่ง' / 'ราคา(+ส่วนปรับ)' คั่นด้วยบรรทัดว่าง"""
    job, err = resolve_job(job_id, date_text)
    if err:
        return f"⚠️ {err}"
    if _job_state(job) == "cancelled":
        return f"⚠️ {job.get('id')} ถูกยกเลิกแล้ว"
    adj = job.get("price_adj")
    price = f"{_effective_price(job)}{adj:+d}" if adj else f"{_effective_price(job)}"
    return f"now  {job.get('car_code', '')}\n\n{job.get('pickup', '-')}-{job.get('dropoff', '-')}\n\n{price}"


_HELP_RE = re.compile(r'^\s*(?:/?help|\?|คำสั่ง(?:ทั้งหมด)?|เมนู|วิธีใช้|ช่วยเหลือ)\s*$', re.IGNORECASE)

HELP_TEXT = """📖 รายการคำสั่งของบอท
(พิมพ์คำสั่งตามช่อง ▶️ แล้วบอทจะตอบกลับในแชท)

━━━━━━━━━━━━━━
1️⃣ ดูตารางสรุป
▶️ Update
▶️ Update 18  /  Update 18/10  /  Update พรุ่งนี้
✅ ได้ตารางสรุปย่อของวันนั้น (ไม่ใส่วัน = วันนี้)

━━━━━━━━━━━━━━
2️⃣ ตัดตาราง
▶️ ตัดตาราง
▶️ ตัดตาราง 18
✅ ได้ใบสรุปเต็มเฉพาะงานที่ยังไม่มีทะเบียน

━━━━━━━━━━━━━━
3️⃣ ดูตารางเต็ม
▶️ ดูตารางเต็ม
▶️ ดูตารางเต็ม 18
✅ ได้ใบสรุปเต็มทุกงานของวันนั้น วงเล็บเป็น 🥶 / 🔥 เหมือนก่อนใส่ทะเบียน

━━━━━━━━━━━━━━
4️⃣ ดูงานเดียว
▶️ ดูงาน F01-18
▶️ ดูงาน F01  (ไม่ใส่วัน = วันนี้)
✅ ได้ใบสรุปเต็มของงานนั้นอย่างเดียว วงเล็บเป็น 🥶 / 🔥

━━━━━━━━━━━━━━
5️⃣ งานด่วน NOW
▶️ now F01
▶️ นาว f01
✅ ได้ข้อความงานด่วน 3 บรรทัด: now ขนาดรถ / จุดรับ-จุดส่ง / ราคา

━━━━━━━━━━━━━━
6️⃣ ปิดใบงาน (ใส่ทะเบียน)
▶️ วิธีที่ 1: ส่งข้อมูลคนขับ (NAME / CAR NO / TEL / APP / TYPE CAR) พร้อมใบงานทั้งใบ จะวางคนขับไว้บนหรือล่างก็ได้ (ทะเบียนเขียนเป็น CAR NO, Car Plate, ทะเบียนรถ หรือ ทะเบียน ก็ได้)
▶️ วิธีที่ 2: บรรทัดแรกพิมพ์  ใบงาน F01  แล้วบรรทัดถัดไปวางข้อมูลคนขับ (บอทดึงใบงานเดิมมาเอง)
✅ ใส่ทะเบียนให้ทันที (ข้อความอื่นก่อน CAR NO เช่น Update ไม่สนใจ) ถ้าส่งมาถูกลำดับอยู่แล้วบอทไม่ตอบกลับ ถ้าสลับลำดับ บอทตอบข้อความที่จัดใหม่ให้: รหัส F ข้อมูลคนขับไว้บน / รหัสอื่น (VK, TT ...) ไว้ล่าง

━━━━━━━━━━━━━━
7️⃣ คืนงาน (ป้าย 📌)
▶️ F22 คืน 16
▶️ F22 คืน 16/10
✅ เปลี่ยนป้ายงาน F22 ของวันที่ 16 เป็น 📌
(รหัสซ้ำหลายวันแล้วไม่ใส่วัน บอทจะถามวัน)

━━━━━━━━━━━━━━
8️⃣ คืนงานด้วยใบงานเต็ม
▶️ บรรทัดแรกพิมพ์  📌คืน
▶️ บรรทัดถัดไปวางใบงานเดิมทั้งใบ
✅ บอทหาจากเลข Order แล้วเปลี่ยนเป็น 📌

━━━━━━━━━━━━━━
9️⃣ ยกเลิกงาน
▶️ วางใบงานเดิม แล้วเพิ่มบรรทัดสุดท้าย  cancel  (หรือ ยกเลิก)
✅ ในตารางจะเป็น  F09 ❌CANCEL❌
▶️ ถ้ายกเลิกผิด พิมพ์  uncancel F09  หรือ  un F09  (ระบุวันได้ เช่น un F09 18)
▶️ หรือวางใบงานเต็ม แล้วเพิ่มบรรทัดสุดท้าย  uncancel  (หรือ  un)  (ใช้เมื่อรหัสซ้ำหลายใบ)
✅ ดึงงานกลับมา ใบสรุปกลับเป็นปกติ (ทะเบียนและราคาที่ปรับไว้เดิมยังอยู่) รายการ CANCEL สั้น ๆ ที่ค้างจะถูกลบ

━━━━━━━━━━━━━━
🔟 แก้ไขราคา
▶️ บรรทัดแรกพิมพ์  แก้ไขราคา  (หรือ อัพเดทราคา)
▶️ บรรทัดถัดไป  F01 +50  /  F02 -30  (มี + หรือ - = เพิ่ม/ลดจากราคาเดิม)
▶️ หรือ  F01 200  (ไม่มี + = ใช้ราคา 200 แทน)
▶️ ระบุวันได้ เช่น  F01 19 +50  /  F01 19 200
✅ ราคาในตารางเป็น 380+50 หรือ 200 (เก็บถาวร, +0 = ยกเลิกส่วนปรับ)

━━━━━━━━━━━━━━
1️⃣1️⃣ แก้ไขใบสรุป
▶️ รหัส [วัน] ค่าเดิม>ค่าใหม่  (พิมพ์หลายบรรทัดได้)
▶️ VK01 19 13:25>20:00  (เวลา)
▶️ VK01 19 แอร์สุ-บางกะปิ>แอร์สุ-สีลม  (เส้นทาง)
▶️ VK01 19 5S>EV-SUV  (ขนาดรถ ราคาตามตารางรถ)
▶️ VK01 19 49VKK4JEQ18>50VKK4JEQ20  (เลข Order)
✅ แก้ในใบสรุป ค่าเดิมต้องตรงกับที่ระบบมี ไม่ใส่วัน = วันนี้

━━━━━━━━━━━━━━
1️⃣2️⃣ ล้างทะเบียน
▶️ ล้างทะเบียน VK05
▶️ ล้างทะเบียน VK05 18
✅ เอาทะเบียน (หรือ 📌) ออก กลับเป็น 🥶 / 🔥

━━━━━━━━━━━━━━
1️⃣3️⃣ ขอ Group ID
▶️ id  (พิมพ์ในกลุ่ม)
✅ บอทตอบรหัสกลุ่ม

━━━━━━━━━━━━━━
1️⃣4️⃣ ดูคำสั่งทั้งหมด
▶️ help  /  คำสั่ง  /  เมนู
✅ แสดงรายการนี้

━━━━━━━━━━━━━━
📝 โยนใบงานเต็มในกลุ่มหรือแชทบอทได้เลย บอทย่อให้แล้วตอบกลับในที่ที่ส่งมา (ใช้ได้ทั้งกลุ่มและ LINE OA)
📝 บอทตอบกลับในแชทที่พิมพ์เท่านั้น ไม่ส่งไปกลุ่มอื่น (ไม่เปลืองโควต้า)
📝 กลุ่มไหนไม่อยากให้บอทอ่านใบงาน ไปเพิ่มกลุ่มในหน้า Settings แล้วเอาติ๊ก "อ่านใบงาน" ออก
📝 ดาวน์โหลดข้อมูลทั้งหมดเป็น Excel หรือส่งไป Google Sheet ได้จากหน้าเว็บ (เมนู "บันทึกใบงาน")"""


def _date_sort_key(d):
    for fmt in ("%d/%m/%Y", "%Y/%m/%d", "%d/%m/%y"):
        try:
            return (0, datetime.datetime.strptime(d, fmt))
        except Exception:
            pass
    return (1, datetime.datetime.max)


def _id_sort_key(job):
    """เรียงตามตัวอักษรแล้วตามตัวเลขของรหัสใบงาน เช่น F01, F02, VK01, VK04"""
    job_id = str(job.get("id", "")).strip().upper()
    m = re.match(r'^([^\d]*)(\d+)', job_id)
    if m:
        return (m.group(1), int(m.group(2)), job_id)
    return (job_id, 10**9, job_id)


def generate_batch_summary(jobs):
    if not jobs:
        return ""

    date_groups = {}
    for job in jobs:
        date_groups.setdefault(_norm_job_date(job.get("date")), []).append(job)

    blocks = []
    sorted_dates = sorted(date_groups.keys(), key=_date_sort_key)

    for d in sorted_dates:
        blocks.append(f"📅 {d}")
        blocks.append("")
        sorted_jobs = sorted(date_groups[d], key=_id_sort_key)
        for i, job in enumerate(sorted_jobs):
            blocks.append(job["formatted_summary"])
            if i < len(sorted_jobs) - 1:
                blocks.append("")
        if d != sorted_dates[-1]:
            blocks.append("")

    return "\n".join(blocks)


def generate_compact_list(jobs):
    """รายการย่อทั้งวัน: บรรทัดแรกของทุกใบงาน แยกตามวันที่ เรียงตามรหัสใบงาน"""
    date_groups = {}
    for job in jobs:
        date_groups.setdefault(_norm_job_date(job.get("date")), []).append(job)
    blocks = []
    for d in sorted(date_groups, key=_date_sort_key):
        blocks.append(f"📅 {d}")
        for job in sorted(date_groups[d], key=_id_sort_key):
            blocks.append(str(job["formatted_summary"]).split("\n")[0].strip())
        blocks.append("")
    return "\n".join(blocks).strip()


def process_batch_jobs():
    """ส่งสรุปใบงานทั้งหมดในคิว: LINE ก่อน แล้วค่อยบันทึก Google Sheets"""
    global job_queue, timer_thread, timer_start_time, last_sender_id, pending_update_dates, last_reply

    with state_lock:
        if timer_thread is not None and threading.current_thread() is not timer_thread:
            timer_thread.cancel()
        jobs = list(job_queue)
        update_dates = list(pending_update_dates)  # วันที่มีการอัปเดตป้ายทะเบียน
        list_dates = list(update_dates)            # วันที่ต้องส่งรายการย่อทั้งวัน
        for j in jobs:
            d = _norm_job_date(j.get("date"))
            if d not in list_dates:
                list_dates.append(d)
        # รายการย่อ: ไม่เอาของวันที่ผ่านมาแล้ว (ตามเวลาไทย) ถ้ามีงานของวันนี้ให้ส่งของพรุ่งนี้ต่อท้ายด้วย
        today = now_th()
        today_str = today.strftime("%d/%m/%Y")
        tomorrow_str = (today + datetime.timedelta(days=1)).strftime("%d/%m/%Y")
        list_dates = [d for d in list_dates if _date_sort_key(d) >= _date_sort_key(today_str)]
        if today_str in list_dates and tomorrow_str not in list_dates:
            list_dates.append(tomorrow_str)
        target_sender_id = last_sender_id
        replies = dict(last_reply)
        last_reply = {}
        # เคลียร์คิวทันที ใบงานใหม่ที่เข้ามาระหว่างส่งจะเริ่มรอบใหม่ได้เลย
        job_queue = []
        pending_update_dates = []
        timer_thread = None
        timer_start_time = None
        last_sender_id = None

        # ใบสรุปปกติ = เฉพาะใบงานใหม่ (อัปเดตป้ายทะเบียนส่งแค่รายการย่อ ไม่ส่งใบสรุปเต็มซ้ำ)
        to_send = list(jobs)
        day_jobs = {}
        for job in latest_jobs:  # ใหม่สุดก่อน ใบงานซ้ำเอาอันล่าสุด
            d = _norm_job_date(job.get("date"))
            if d not in list_dates or job.get("status") == "ยกเลิก":
                continue
            key = (d, _job_key(job))
            if key in day_jobs:
                continue
            day_jobs[key] = job
        all_day_jobs = list(day_jobs.values())

    print(f"🚀 เริ่มส่งสรุป: ใบงานใหม่ {len(jobs)} / รายการย่อ {len(all_day_jobs)}")
    if not jobs and not all_day_jobs:
        persist_state()
        return

    try:
        summary_text = generate_batch_summary(to_send)
        list_text = generate_compact_list(all_day_jobs)
        if summary_text or list_text:
            # LINE จำกัด 5000 ตัวอักษรต่อข้อความ
            chunks = [summary_text[i:i + 4900] for i in range(0, len(summary_text), 4900)][:4]
            # ข้อความเพิ่มเติม: รายการย่อของใบงานทั้งวัน (เฉพาะบรรทัดแรกของแต่ละใบงาน)
            if list_text:
                chunks.append(list_text[:4900])
            results = []

            if not PUSH_ENABLED:
                # โหมดประหยัดโควต้า: ตอบกลับ (reply) เฉพาะแชทที่พิมพ์มา ไม่ส่งเข้ากลุ่มอื่น ไม่ push
                if not replies:
                    print("ℹ️ โหมดตอบกลับอย่างเดียว: ไม่มีแชทให้ตอบ (ไม่ push) ดูผลได้ด้วยคำสั่ง Update")
                for chat_id, rep in replies.items():
                    label = f"แชท {str(chat_id)[:8]}…"
                    try:
                        with ApiClient(rep.get("conf") or configuration) as api_client:
                            MessagingApi(api_client).reply_message(ReplyMessageRequest(
                                reply_token=rep["token"], messages=[TextMessage(text=c) for c in chunks]))
                        print(f"✅ ตอบกลับ {label} สำเร็จ")
                        results.append((label, True, "reply"))
                    except Exception as e:
                        print(f"❌ ตอบกลับ {label} ไม่ได้ (reply token อาจหมดอายุ): {e}")
                        results.append((label, False, "reply"))
            else:
                settings = load_settings()
                active_groups = [g for g in settings.get("line_groups", [])
                                 if g.get("group_id") and g.get("send", True)]
                targets = []
                if target_sender_id:
                    targets.append(("ผู้ส่ง (LINE OA)", target_sender_id))
                for grp in active_groups:
                    targets.append((f"กลุ่ม {grp.get('name')}", grp.get("group_id")))
                print(f"📤 ส่งใบสรุปไปยัง: {[(l, str(t)[:8] + '…') for l, t in targets]} "
                      f"(กลุ่มที่ติ๊กส่ง {len(active_groups)} จาก {len(settings.get('line_groups', []))})")
                for label, to in targets:
                    messages = [TextMessage(text=c) for c in chunks]
                    # แชทที่ส่งใบงานมา: ตอบด้วย reply ก่อน (ฟรี) ถ้าไม่ได้ค่อยใช้ push
                    rep = replies.pop(to, None)
                    if rep:
                        try:
                            with ApiClient(rep.get("conf") or configuration) as api_client:
                                MessagingApi(api_client).reply_message(
                                    ReplyMessageRequest(reply_token=rep["token"], messages=messages))
                            results.append((label, True, "reply"))
                            continue
                        except Exception as e:
                            print(f"⚠️ ตอบกลับ {label} ด้วย reply ไม่ได้ ลองใช้ push แทน: {e}")
                    results.append((label, push_with_fallback(to, messages, label), "push"))

            if summary_text:
                log_sent("summary", summary_text, sorted({_norm_job_date(j.get("date")) for j in to_send}), results)
            if list_text:
                log_sent("compact", list_text, list_dates, results)

            if GOOGLE_SPREADSHEET_ID and jobs:
                try:
                    ok = save_to_google_sheets(GOOGLE_SPREADSHEET_ID, [j["text"] for j in jobs], jobs)
                    print("✅ บันทึกข้อมูลลง Google Sheets สำเร็จ" if ok else "❌ บันทึก Google Sheets ล้มเหลว")
                except Exception as e:
                    print(f"❌ บันทึก Google Sheets ล้มเหลว: {e}")
    except Exception as e:
        print(f"❌ ประมวลผลคิวใบงานผิดพลาด: {e}")
    finally:
        sent_keys = {(_job_key(j), j.get("id")) for j in jobs}
        with state_lock:
            for job in latest_jobs:
                if (_job_key(job), job.get("id")) in sent_keys:
                    job["status"] = "ส่งแล้ว"
            _prune_jobs()
        persist_state()


@handler.add(JoinEvent)
def handle_join(event):
    src = event.source
    cid = getattr(src, 'group_id', None) or getattr(src, 'room_id', None)
    print(f"➕ บอทถูกเชิญเข้า {src.type}: {cid}")
    _record_group(cid, "main")


@handler.add(LeaveEvent)
def handle_leave(event):
    src = event.source
    cid = getattr(src, 'group_id', None) or getattr(src, 'room_id', None)
    print(f"➖ บอทถูกนำออกจาก {src.type}: {cid}")


@handler.add(MessageEvent, message=TextMessageContent)
def handle_message(event):
    _handle_text(event, _pick_conf(configuration, "บอทหลัก"))


def _handle_text(event, conf):
    """จัดการข้อความที่เข้ามา conf = บอทที่รับข้อความ (บอทหลักหรือบอทสำรอง) ใช้ตอบกลับด้วยบอทตัวเดียวกัน"""
    received_text = event.message.text.strip()
    source_type = event.source.type
    user_id = getattr(event.source, 'user_id', None)
    first_line = received_text.splitlines()[0][:40] if received_text else ""
    print(f"📩 ได้รับข้อความจาก {source_type}: {first_line!r}")

    bot_key = _bot_key_of(conf)
    ctx_group = getattr(event.source, 'group_id', None) or getattr(event.source, 'room_id', None)
    _record_group(ctx_group, bot_key)

    def allow(cmd, notify=True):
        """เช็คสิทธิ์คำสั่ง (บอท + กลุ่ม) ที่ตั้งไว้ในหน้าเว็บ ถ้าปิดไว้ตอบแจ้งสั้น ๆ (ยกเว้น notify=False)"""
        if _cmd_allowed(cmd, bot_key, ctx_group):
            return True
        print(f"⛔ ปิดคำสั่ง {cmd} ({bot_key}{', กลุ่ม ' + str(ctx_group) if ctx_group else ''})")
        if notify and not ctx_group:  # ในกลุ่มเงียบ ไม่ตอบอะไร ส่วนแชทส่วนตัวตอบแจ้งสั้น ๆ
            _reply_text(conf, event.reply_token, "⛔ คำสั่งนี้ถูกปิดใช้งานสำหรับบอทนี้")
        return False

    if source_type == 'group' and received_text.lower() == "id":
        if not allow("group_id"):
            return
        print(f"🆔 ตอบ ID ของกลุ่ม: {event.source.group_id}")
        _reply_text(conf, event.reply_token, f"{event.source.group_id}")
        return

    # "Update 17/10" -> ตอบเฉพาะตารางสรุปย่อของวันนั้น (ใช้ reply ไม่เสียโควต้า ไม่ส่งเข้ากลุ่มอื่น)
    um = _UPDATE_RE.match(received_text)
    if um:
        if not allow("update"):
            return
        date = parse_update_date(um.group(1))
        if date is None:
            reply = "รูปแบบ: Update 17 หรือ Update 17/10 หรือ Update 17/10/2026 (พิมพ์ Update เฉย ๆ = วันนี้)"
        else:
            reply = generate_compact_list(sorted(jobs_for_date(date), key=_id_sort_key)) or f"ไม่มีใบงานวันที่ {date}"
        print(f"📋 Update {date or um.group(1)!r}")
        _reply_long(conf, event.reply_token, reply)
        return

    # คำสั่งอ่านอย่างเดียว (ใช้ได้ทุกกลุ่มที่บอทอยู่): help / ตัดตาราง / NOW
    if _HELP_RE.match(received_text):
        if not allow("help"):
            return
        print("📖 help")
        _reply_long(conf, event.reply_token, HELP_TEXT)
        return
    fm = _FULL_TABLE_RE.match(received_text)
    if fm:
        if not allow("full_table"):
            return
        print(f"📋 ดูตารางเต็ม {fm.group(1)!r}")
        _reply_long(conf, event.reply_token, full_table_text(fm.group(1)))
        return
    vm = _VIEW_JOB_RE.match(received_text)
    if vm:
        if not allow("view_job"):
            return
        print(f"🔎 ดูงาน {vm.group(1)}")
        _reply_long(conf, event.reply_token, view_job_text(vm.group(1), vm.group(2)))
        return
    cm = _CUT_RE.match(received_text)
    if cm:
        if not allow("cut_table"):
            return
        print(f"✂️ ตัดตาราง {cm.group(1)!r}")
        _reply_long(conf, event.reply_token, cut_table_text(cm.group(1)))
        return
    nm = _NOW_RE.match(received_text)
    if nm:
        if not allow("now"):
            return
        print(f"‼️ NOW {nm.group(1)}")
        _reply_long(conf, event.reply_token, now_text(nm.group(1), nm.group(2)))
        return

    # อ่านใบงานในทุกกลุ่มที่บอทอยู่ ยกเว้นกลุ่มที่ตั้งไว้ในหน้าเว็บแล้วเอาติ๊ก "อ่านใบงาน" ออก
    if source_type in ('group', 'room'):
        chat_id = getattr(event.source, 'group_id', None) or getattr(event.source, 'room_id', None)
        grp = next((g for g in load_settings().get("line_groups", [])
                    if str(g.get("group_id", "")).strip() == chat_id), None)
        if grp is not None and not grp.get("read", True):
            print(f"🚫 ไม่อ่าน: กลุ่ม {chat_id} ปิด 'อ่านใบงาน' ไว้ในหน้าเว็บ")
            return

    chat_id = (getattr(event.source, 'group_id', None) or getattr(event.source, 'room_id', None)
               or user_id)

    # "📌คืน" + ใบงานเดิมทั้งใบ -> ค้นใบงานจากเลข Order แล้วเปลี่ยนป้ายเป็น 📌
    rm = _RETURN_TICKET_RE.match(received_text)
    if rm:
        if not allow("return"):
            return
        ok, result = update_plate(rm.group(1).strip(), RETURN_MARK, sender_id=user_id)
        print(f"📌 คืนงานจากใบงาน {result['id']}" if ok else f"⚠️ คืนงานไม่สำเร็จ: {result}")
        if ok:
            _remember_reply(event.reply_token, chat_id, conf)
        return

    ret = parse_return_command(received_text)
    if ret:
        if not allow("return"):
            return
        ok, result = return_plate(*ret, sender_id=user_id)
        if ok:
            print(f"📌 คืนป้าย {result['id']} วันที่ {result['date']}")
            _remember_reply(event.reply_token, chat_id, conf)
        else:
            print(f"⚠️ {result}")
            try:
                with ApiClient(conf) as api_client:
                    MessagingApi(api_client).reply_message(ReplyMessageRequest(
                        reply_token=event.reply_token, messages=[TextMessage(text=f"⚠️ {result}")]))
            except Exception as e:
                print(f"⚠️ ตอบกลับไม่ได้: {e}")
        return

    um2 = _UNCANCEL_RE.match(received_text)
    if um2:
        if not allow("uncancel"):
            return
        ok, result = uncancel_job(um2.group(1), um2.group(2), user_id)
        print(f"↩️ uncancel {um2.group(1)}" if ok else f"⚠️ uncancel ไม่สำเร็จ: {result}")
        _reply_long(conf, event.reply_token, result if ok else f"⚠️ {result}")
        return
    if _UNCANCEL_LINE_RE.search(received_text) and _is_ticket_text(received_text):
        if not allow("uncancel"):
            return
        ok, result = uncancel_ticket(received_text, user_id)
        print("↩️ uncancel (ใบงาน)" if ok else f"⚠️ uncancel ไม่สำเร็จ: {result}")
        _reply_long(conf, event.reply_token, result if ok else f"⚠️ {result}")
        return
    if is_cancel_ticket(received_text) and (
            "【客户订单号】" in received_text or is_voucher_format(received_text)
            or is_kv_format(received_text) or is_plain_format(received_text)):
        if not allow("cancel"):
            return
        ok, result = cancel_job(received_text, sender_id=user_id)
        print(f"❌ ยกเลิกงาน {result['id']}" if ok else f"⚠️ ยกเลิกงานไม่สำเร็จ: {result}")
        if ok:
            _remember_reply(event.reply_token, chat_id, conf)
        return

    # คำสั่งแก้ข้อมูล: แก้ไขราคา / ปิดใบงาน
    msg_lines = received_text.splitlines()
    first_line_text = msg_lines[0] if msg_lines else ""
    cp = _CLEAR_PLATE_RE.match(received_text)
    if cp:
        if not allow("clear_plate"):
            return
        ok, result = clear_plate(cp.group(1), cp.group(2), user_id)
        print(f"🧹 ล้างทะเบียน {cp.group(1)}" if ok else f"⚠️ ล้างทะเบียนไม่สำเร็จ: {result}")
        _reply_long(conf, event.reply_token, result if ok else f"⚠️ {result}")
        return
    if is_edit_message(msg_lines):
        if not allow("edit_job"):
            return
        print("✏️ แก้ไขใบสรุป")
        _reply_long(conf, event.reply_token, edit_job_command(msg_lines, user_id))
        return
    if _PRICE_HEAD_RE.match(first_line_text):
        if not allow("price"):
            return
        print("💰 แก้ไขราคา")
        _reply_long(conf, event.reply_token, price_command(msg_lines[1:], user_id))
        return
    cl = _CLOSE_RE.match(first_line_text)
    if cl and len(msg_lines) > 1 and extract_plate(received_text):
        if not allow("close_job"):
            return
        ok, result = close_job(cl.group(1), cl.group(2), "\n".join(msg_lines[1:]), user_id)
        print(f"🔒 ปิดใบงาน {cl.group(1)}" if ok else f"⚠️ ปิดใบงานไม่สำเร็จ: {result}")
        _reply_long(conf, event.reply_token, result if ok else f"⚠️ {result}")
        return
    if extract_plate(received_text):
        if not allow("close_job"):
            return
        ok, result = close_with_ticket(received_text, sender_id=user_id)
        print(("🔒 ปิดใบงาน (ใบงาน + ข้อมูลคนขับ)" + ("" if result else " ลำดับถูกอยู่แล้ว ไม่ตอบกลับ")) if ok
              else f"⚠️ ปิดใบงานไม่สำเร็จ: {result}")
        if not ok:
            _reply_long(conf, event.reply_token, f"⚠️ {result}")
        elif result:
            _reply_long(conf, event.reply_token, result)
        return

    if ("【客户订单号】" not in received_text and not is_voucher_format(received_text)
            and not is_kv_format(received_text) and not is_plain_format(received_text)):
        print("ℹ️ ไม่ใช่ใบงาน (ไม่มี 【客户订单号】 หรือบรรทัด ชูป้าย)")
        return

    if not allow("ticket", notify=False):
        return
    ok, result = add_job_to_queue(received_text, sender_id=user_id)
    print(f"✅ รับใบงาน {result['id']} เข้าคิว" if ok else f"⚠️ ไม่รับใบงาน: {result}")
    if ok:
        _remember_reply(event.reply_token, chat_id, conf)
    else:
        _reply_long(conf, event.reply_token, str(result))


def _remember_reply(token, chat_id, conf=None):
    """เก็บ reply token ของข้อความล่าสุด ไว้ตอบใบสรุปกลับแชทนั้นโดยไม่เสียโควต้า"""
    global last_reply
    if token and chat_id:
        with state_lock:
            last_reply[chat_id] = {"token": token, "conf": conf or configuration}


# ====================== หน้าเว็บ / API ======================
@app.route("/")
def index():
    return render_template("ui.html")


@app.route("/health")
def health():
    return "OK", 200


@app.route("/api/status", methods=["GET"])
def api_status():
    settings = load_settings()
    wait_seconds = int(settings.get("wait_seconds", 10))

    with state_lock:
        time_left = 0
        if timer_start_time:
            time_left = max(0, int(wait_seconds - (time.time() - timer_start_time)))
        queue_count = len(job_queue)
        jobs_copy = [dict(j, uid=_job_uid(j)) for j in latest_jobs]

    return jsonify({
        "line_status": bool(LINE_CHANNEL_ACCESS_TOKEN and LINE_CHANNEL_SECRET),
        "sheets_status": bool(GOOGLE_SPREADSHEET_ID and has_credentials()),
        "storage_ok": not _settings_error,
        "storage_mode": "supabase" if storage.using_supabase() else "local",
        "queue_count": queue_count,
        "time_left_seconds": time_left,
        "max_wait_seconds": wait_seconds,
        "latest_jobs": jobs_copy,
        "today_stats": _today_stats(),
        "custom_keywords": settings.get("custom_keywords", []),
        "line_groups": settings.get("line_groups", []),
        "main_roads_dict": settings.get("main_roads_dict", DEFAULT_MAIN_ROADS_DICT),
        "major_areas_dict": settings.get("major_areas_dict", DEFAULT_MAJOR_AREAS_DICT),
        "settings": settings
    })


@app.route("/api/records", methods=["GET"])
def api_records():
    """ข้อมูลใบงานแบบเป็นระบบ แยกตามวันที่: ใบเต็ม / ใบย่อ / คืนงาน / ยกเลิกงาน / ประวัติเหตุการณ์ / ประวัติการส่ง"""
    today = now_th().strftime("%d/%m/%Y")
    date = _norm_job_date(request.args.get("date") or today)
    with state_lock:
        jobs, seen = [], set()
        for j in latest_jobs:  # ใหม่สุดก่อน ใบซ้ำเอาอันล่าสุด
            if j.get("status") == "ยกเลิก":
                continue
            item = dict(j)
            item["date"] = _norm_job_date(j.get("date"))
            k = (item["date"], _job_key(j))
            if k in seen:
                continue
            seen.add(k)
            item["state"] = _job_state(j)
            item["uid"] = _job_uid(j)
            jobs.append(item)
        activity, sent = list(_activity_log), list(_sent_log)

    dates = sorted({j["date"] for j in jobs} | {today} | set(history_dates()), key=_date_sort_key, reverse=True)
    full = sorted([j for j in jobs if j["date"] == date], key=_id_sort_key)
    if not full and storage.using_supabase() and _iso_date(date):
        try:
            rows = storage.select_rows("jobs", {"job_date": f"eq.{_iso_date(date)}", "select": "*", "limit": "1000"})
            full = sorted(_rows_to_jobs(rows), key=_id_sort_key)
        except Exception as e:
            print(f"⚠️ อ่านตาราง jobs ไม่สำเร็จ: {e}")
    return jsonify({
        "date": date, "today": today, "dates": dates,
        "full": full,
        "compact": generate_compact_list(full) if full else "",
        "compact_uids": [j.get("uid") or _job_uid(j) for j in full],
        "returns": [j for j in full if j["state"] == "returned"],
        "cancels": [j for j in full if j["state"] == "cancelled"],
        "activity": [a for a in reversed(activity) if a.get("date") == date][:200],
        "sent": [x for x in reversed(sent) if date in x.get("dates", [])][:50],
    })


def _today_stats():
    today = now_th().strftime("%d/%m/%Y")
    with state_lock:
        seen, rows = set(), []
        for j in latest_jobs:
            if j.get("status") == "ยกเลิก" or _norm_job_date(j.get("date")) != today:
                continue
            k = _job_key(j)
            if k in seen:
                continue
            seen.add(k)
            rows.append(j)
    states = [_job_state(j) for j in rows]
    return {
        "date": today, "total": len(rows),
        "sent": sum(1 for j in rows if j.get("status") == "ส่งแล้ว"),
        "waiting": sum(1 for j in rows if j.get("status") == "รอส่ง"),
        "returned": states.count("returned"), "cancelled": states.count("cancelled"),
    }


_sheets_health_cache = {"at": 0, "ok": False}


@app.route("/api/system_health", methods=["GET"])
def api_system_health():
    line_connected = bool(LINE_CHANNEL_ACCESS_TOKEN and LINE_CHANNEL_SECRET)

    # เช็ค Google Sheets ไม่เกินทุก 5 นาที (กันโดนจำกัดการใช้งาน)
    if GOOGLE_SPREADSHEET_ID and time.time() - _sheets_health_cache["at"] > 300:
        ok = False
        try:
            client = get_sheet_client()
            if client:
                client.open_by_key(GOOGLE_SPREADSHEET_ID)
                ok = True
        except Exception as e:
            print(f"Google Sheets check error: {e}")
        _sheets_health_cache.update(at=time.time(), ok=ok)

    storage_ok, storage_msg = storage.health()

    return jsonify({
        "line_bot": line_connected,
        "google_sheets": bool(GOOGLE_SPREADSHEET_ID) and _sheets_health_cache["ok"],
        "storage": storage_ok,
        "storage_message": storage_msg,
        "bots": bots_status(),
    })


@app.route("/callback", methods=['GET', 'POST'])
def callback():
    if request.method == 'GET':
        return 'OK', 200

    signature = request.headers.get('X-Line-Signature', '')
    body = request.get_data(as_text=True)
    print("📨 ได้รับ webhook จาก LINE")
    try:
        _ctx.destination = json.loads(body).get("destination")
    except Exception:
        _ctx.destination = None

    try:
        handler.handle(body, signature)
        return 'OK'
    except InvalidSignatureError:
        pass
    # ลายเซ็นไม่ตรงบอทหลัก: ลองบอทสำรอง (ใช้ตอบคำสั่ง id ในกลุ่ม)
    for bh in BACKUP_HANDLERS:
        try:
            bh.handle(body, signature)
            return 'OK'
        except InvalidSignatureError:
            continue
    print(f"❌ ลายเซ็น LINE ไม่ถูกต้อง: ลองแล้ว {1 + len(BACKUP_HANDLERS)} บอท (บอทหลัก + บอทสำรองที่มี secret {len(BACKUP_HANDLERS)} ตัว) "
          "ไม่มี secret ตัวไหนตรงกับบอทที่ส่งมา เช็ค LINE_BACKUP_SECRETS ว่าใส่ Channel secret ของบอทนี้ ไม่ใช่ access token และเรียงลำดับตรงกับ LINE_BACKUP_TOKENS")
    abort(400)


@app.route("/api/get_sheets_data", methods=["GET"])
def api_get_sheets_data():
    client = get_sheet_client()
    if not client or not GOOGLE_SPREADSHEET_ID:
        return jsonify({"success": False, "message": "Google Sheets not connected"}), 400

    try:
        spreadsheet = client.open_by_key(GOOGLE_SPREADSHEET_ID)
        rows = spreadsheet.worksheet("SUMMARY").get_all_records()
        return jsonify({"success": True, "data": rows})
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/api/add_job", methods=["POST"])
def api_add_job():
    data = request.get_json(silent=True) or {}
    text = str(data.get("text", "")).strip()
    if not text:
        return jsonify({"success": False, "message": "ข้อความว่างเปล่า"}), 400
    plate = extract_plate(text)
    ok, result = update_plate(text, plate) if plate else add_job_to_queue(text)
    if ok:
        return jsonify({"success": True, "id": result["id"], "message": "เพิ่มใบงานสำเร็จ"})
    return jsonify({"success": False, "message": result})


@app.route("/api/save_settings", methods=["POST"])
def api_save_settings():
    new_settings = request.get_json(silent=True) or {}
    ok, msg = save_settings(new_settings)
    return jsonify({"success": ok, "message": msg}), (200 if ok else 500)


@app.route("/api/trigger_send", methods=["POST"])
def api_trigger_send():
    process_batch_jobs()
    return jsonify({"success": True, "message": "ส่งสรุปใบงานเรียบร้อยแล้ว"})


@app.route('/api/delete_job', methods=['POST'])
@app.route("/api/delete_job/<job_id>", methods=["DELETE", "POST"])
def api_delete_job(job_id=None):
    try:
        if not job_id:
            data = request.get_json(silent=True) or {}
            job_id = data.get('id')

        if not job_id:
            return jsonify({"success": False, "message": "ไม่พบรหัสใบงานที่ต้องการลบ"}), 400

        if not GOOGLE_SPREADSHEET_ID:
            return jsonify({"success": False, "message": "ยังไม่ได้ตั้งค่า GOOGLE_SPREADSHEET_ID"}), 400

        if delete_row_from_google_sheets(GOOGLE_SPREADSHEET_ID, job_id):
            return jsonify({"success": True, "message": f"ลบใบงาน {job_id} สำเร็จ"})
        return jsonify({"success": False, "message": "ไม่พบข้อมูลใน Google Sheets หรือเกิดข้อผิดพลาด"}), 404

    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/api/delete_record", methods=["POST"])
def api_delete_record():
    """ลบใบงานออกจากระบบถาวร (หน่วยความจำ + คิว + ตาราง jobs ใน Supabase) อ้างด้วย uid จากหน้าเว็บ"""
    global job_queue, latest_jobs
    data = request.get_json(silent=True) or {}
    uid = str(data.get("uid", "")).strip()
    if not uid:
        return jsonify({"success": False, "message": "ไม่พบรหัสใบงาน (uid)"}), 400
    with state_lock:
        victims = [j for j in latest_jobs + job_queue if _job_uid(j) == uid]
        ids = {id(j) for j in victims}
        latest_jobs = [j for j in latest_jobs if id(j) not in ids]
        job_queue = [j for j in job_queue if id(j) not in ids]
        sample = victims[0] if victims else None
        if sample:
            log_activity("delete", sample, "ลบจากหน้าเว็บ")
    in_table = False
    try:
        in_table = bool(storage.delete_rows("jobs", "uid", uid))
    except Exception as e:
        print(f"⚠️ ลบแถวในตาราง jobs ไม่สำเร็จ: {e}")
    _jobs_synced.pop(uid, None)
    if not victims and not in_table:
        return jsonify({"success": False, "message": "ไม่พบใบงานนี้ (อาจถูกลบไปแล้ว)"}), 404
    persist_state()
    return jsonify({"success": True, "message": f"ลบใบงาน {sample.get('id') if sample else ''} เรียบร้อยแล้ว"})


@app.route("/api/update_compact", methods=["POST"])
def api_update_compact():
    """แก้ไขใบงานย่อจากหน้าเว็บ: 1 บรรทัด = 1 ใบงาน (ตามลำดับ uids ที่หน้าเว็บส่งมา) แล้วบันทึกลงบรรทัดแรกของใบสรุปนั้น"""
    data = request.get_json(silent=True) or {}
    uids = data.get("uids") or []
    lines = [str(x).strip() for x in (data.get("lines") or [])]
    if not uids or len(uids) != len(lines):
        return jsonify({"success": False, "message": f"จำนวนบรรทัดไม่ตรงกับจำนวนใบงาน (ต้องมี {len(uids)} บรรทัด ห้ามเพิ่ม/ลบบรรทัด)"}), 400
    if any(not ln for ln in lines):
        return jsonify({"success": False, "message": "มีบรรทัดว่าง กรุณากรอกให้ครบทุกบรรทัด"}), 400

    def with_first_line(summary, new_line):
        parts = str(summary or "").split("\n")
        parts[0] = new_line
        return "\n".join(parts)

    changed, missing = 0, []
    with state_lock:
        for uid, new_line in zip(uids, lines):
            hits = [j for j in latest_jobs + job_queue if _job_uid(j) == uid]
            if not hits:
                missing.append((uid, new_line))
                continue
            for j in hits:
                old_line = str(j.get("formatted_summary", "")).split("\n")[0].strip()
                if old_line == new_line:
                    continue
                j["formatted_summary"] = with_first_line(j.get("formatted_summary"), new_line)
                mark_job(j, "amend", f"แก้ใบงานย่อจากหน้าเว็บ: {old_line} > {new_line}")
                changed += 1

    # ใบงานเก่าที่อยู่เฉพาะในตาราง jobs (ไม่อยู่ในหน่วยความจำ)
    for uid, new_line in missing:
        try:
            rows = storage.select_rows("jobs", {"uid": f"eq.{uid}", "select": "*", "limit": "1"})
            if not rows:
                continue
            row = rows[0]
            if str(row.get("summary", "")).split("\n")[0].strip() == new_line:
                continue
            row["summary"] = with_first_line(row.get("summary"), new_line)
            row["events"] = (row.get("events") or []) + [{"ts": _ts(), "type": "amend", "detail": "แก้ใบงานย่อจากหน้าเว็บ"}]
            storage.upsert_rows("jobs", [row], "uid")
            changed += 1
        except Exception as e:
            print(f"⚠️ แก้ใบงานย่อในตาราง jobs ไม่สำเร็จ ({uid}): {e}")
            return jsonify({"success": False, "message": f"บันทึกไม่สำเร็จ: {e}"}), 500

    persist_state()
    return jsonify({"success": True, "changed": changed,
                    "message": f"บันทึกใบงานย่อเรียบร้อยแล้ว ({changed} รายการ)" if changed else "ไม่มีการเปลี่ยนแปลง"})


@app.route("/api/clear_queue", methods=["POST"])
def api_clear_queue():
    global job_queue, timer_thread, timer_start_time, pending_update_dates
    with state_lock:
        if timer_thread:
            timer_thread.cancel()
        num_cleared = len(job_queue)
        pending_update_dates = []
        cleared_keys = {_job_key(j) for j in job_queue}
        job_queue = []
        timer_thread = None
        timer_start_time = None
        for job in latest_jobs:
            if job.get("status") == "รอส่ง" and _job_key(job) in cleared_keys:
                job["status"] = "ยกเลิก"
    persist_state()
    return jsonify({"success": True, "num_cleared": num_cleared, "message": "ล้าง Queue เรียบร้อยแล้ว"})


# ====================== ส่งออกข้อมูล (Excel / Google Sheet) ======================
EXPORT_HEADERS = ["วันที่", "รหัส", "เวลา", "จุดรับ", "จุดส่ง", "เที่ยวบิน", "Order", "รถ", "ราคา", "ปรับราคา",
                  "ทะเบียน", "สถานะส่ง", "สถานะงาน", "ใบสรุป", "ข้อความต้นฉบับ"]
_STATE_TH = {"normal": "ปกติ", "returned": "คืนงาน", "cancelled": "ยกเลิก"}


def _parse_export_date(text):
    """รับ dd/mm/yyyy หรือ yyyy-mm-dd (จากช่องวันที่ของเบราว์เซอร์) คืน dd/mm/yyyy หรือ None"""
    text = (text or "").strip()
    if not text:
        return None
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})$', text)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    return parse_update_date(text)


def collect_export_jobs(date_from=None, date_to=None):
    """ใบงานทั้งหมด (ตาราง jobs ใน Supabase + ที่อยู่ในระบบ) ไม่ซ้ำ เรียงวันที่แล้วรหัส"""
    by_uid = {}
    if storage.using_supabase():
        params = {"select": "*", "order": "job_date.asc", "limit": "20000"}
        a, b = _iso_date(date_from) if date_from else None, _iso_date(date_to) if date_to else None
        if a and b:
            params["and"] = f"(job_date.gte.{a},job_date.lte.{b})"
        elif a:
            params["job_date"] = f"gte.{a}"
        elif b:
            params["job_date"] = f"lte.{b}"
        try:
            for j in _rows_to_jobs(storage.select_rows("jobs", params)):
                by_uid[f"{_norm_job_date(j.get('date'))}|{_job_key(j)}"] = j
        except Exception as e:
            print(f"⚠️ อ่านตาราง jobs เพื่อส่งออกไม่สำเร็จ (ใช้เฉพาะข้อมูลในระบบ): {e}")
    with state_lock:
        for j in reversed(latest_jobs):  # ข้อมูลในระบบใหม่กว่า ทับของในตาราง
            if j.get("status") == "ยกเลิก":
                continue
            item = dict(j)
            item["state"] = _job_state(j)
            by_uid[f"{_norm_job_date(j.get('date'))}|{_job_key(j)}"] = item
    lo = _date_sort_key(date_from) if date_from else None
    hi = _date_sort_key(date_to) if date_to else None
    jobs = []
    for j in by_uid.values():
        d = _date_sort_key(_norm_job_date(j.get("date")))
        if (lo and d < lo) or (hi and d > hi):
            continue
        jobs.append(j)
    jobs.sort(key=lambda j: (_date_sort_key(_norm_job_date(j.get("date"))), _id_sort_key(j)))
    return jobs


def _job_export_row(j):
    adj = j.get("price_adj")
    return [
        _norm_job_date(j.get("date")), j.get("id", ""), j.get("time", ""),
        j.get("pickup", ""), j.get("dropoff", ""), j.get("flight", ""), j.get("order", ""),
        j.get("car_code", ""), _effective_price(j), f"{adj:+d}" if adj else "",
        j.get("plate", "") or "", j.get("status", ""), _STATE_TH.get(_job_state(j), ""),
        j.get("formatted_summary", ""), j.get("text") or j.get("raw_text") or "",
    ]


def _compact_rows(jobs):
    return [[_norm_job_date(j.get("date")), j.get("id", ""), str(j.get("formatted_summary", "")).split("\n")[0]]
            for j in jobs if j.get("status") != "ยกเลิก"]


def _export_tables(jobs):
    """ชื่อชีต -> (หัวตาราง, แถว) ใช้ทั้ง Excel และ Google Sheet"""
    dates = {_norm_job_date(j.get("date")) for j in jobs}
    with state_lock:
        activity = [a for a in _activity_log if not a.get("date") or a["date"] in dates]
        sent = [x for x in _sent_log if any(d in dates for d in x.get("dates", []))]
    rows = [_job_export_row(j) for j in jobs]
    returned = [r for r, j in zip(rows, jobs) if _job_state(j) == "returned"]
    cancelled = [r for r, j in zip(rows, jobs) if _job_state(j) == "cancelled"]
    return {
        "ใบงานทั้งหมด": (EXPORT_HEADERS, rows),
        "ใบงานย่อ": (["วันที่", "รหัส", "ใบงานย่อ"], _compact_rows(jobs)),
        "คืนงาน": (EXPORT_HEADERS, returned),
        "ยกเลิก": (EXPORT_HEADERS, cancelled),
        "บันทึกเหตุการณ์": (["เวลา", "เหตุการณ์", "รหัส", "วันที่", "Order", "รายละเอียด"],
                            [[a.get("ts", ""), a.get("type", ""), a.get("id", ""), a.get("date", ""),
                              a.get("order", ""), a.get("detail", "")] for a in activity]),
        "ประวัติการส่ง": (["เวลา", "ชนิด", "วันที่ของใบงาน", "ปลายทาง", "ข้อความ"],
                          [[x.get("ts", ""), "ใบงานย่อ" if x.get("kind") == "compact" else "ใบสรุป",
                            ", ".join(x.get("dates", [])),
                            ", ".join(f"{t['to']}({'ok' if t['ok'] else 'fail'})" for t in x.get("targets", [])),
                            x.get("text", "")] for x in sent]),
    }


def build_export_xlsx(jobs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in _export_tables(jobs).items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append([str(c) if c is not None else "" for c in r])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        for idx, h in enumerate(headers, start=1):
            longest = max([len(str(h))] + [len(str(r[idx - 1]).split("\n")[0]) for r in rows[:300]]) if rows else len(str(h))
            ws.column_dimensions[get_column_letter(idx)].width = min(max(10, longest + 2), 60)
        if rows:
            ws.auto_filter.ref = ws.dimensions
    import io
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


@app.route("/api/export.xlsx", methods=["GET"])
def api_export_xlsx():
    """ดาวน์โหลดใบงานเป็น Excel: ?from=...&to=... (ไม่ใส่ = ทั้งหมด) รับ dd/mm/yyyy หรือ yyyy-mm-dd"""
    date_from = _parse_export_date(request.args.get("from"))
    date_to = _parse_export_date(request.args.get("to"))
    jobs = collect_export_jobs(date_from, date_to)
    tag = (date_from or "all").replace("/", "-") + (f"_to_{date_to.replace('/', '-')}" if date_to and date_to != date_from else "")
    try:
        buf = build_export_xlsx(jobs)
    except ImportError:
        import csv, io
        out = io.StringIO()
        w = csv.writer(out)
        w.writerow(EXPORT_HEADERS)
        w.writerows(_job_export_row(j) for j in jobs)
        data = io.BytesIO(("﻿" + out.getvalue()).encode("utf-8"))
        return send_file(data, mimetype="text/csv", as_attachment=True, download_name=f"jobs_{tag}.csv")
    return send_file(buf, as_attachment=True, download_name=f"jobs_{tag}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.route("/api/export_sheets", methods=["POST"])
def api_export_sheets():
    """ส่งออกไป Google Sheet (สเปรดชีตเดียวกับที่บันทึกใบงาน) เป็นชีตใหม่ EXPORT_xxx ล้างแล้วเขียนใหม่ทุกครั้ง"""
    if not GOOGLE_SPREADSHEET_ID:
        return jsonify({"success": False, "message": "ยังไม่ได้ตั้งค่า GOOGLE_SPREADSHEET_ID"}), 400
    client = get_sheet_client()
    if not client:
        return jsonify({"success": False, "message": "เชื่อมต่อ Google Sheets ไม่ได้ (เช็ค credentials.json)"}), 400
    data = request.get_json(silent=True) or {}
    jobs = collect_export_jobs(_parse_export_date(data.get("from")), _parse_export_date(data.get("to")))
    try:
        spreadsheet = client.open_by_key(GOOGLE_SPREADSHEET_ID)
        for name, (headers, rows) in _export_tables(jobs).items():
            title = f"EXPORT_{name}"
            try:
                ws = spreadsheet.worksheet(title)
                ws.clear()
            except Exception:
                ws = spreadsheet.add_worksheet(title=title, rows=max(100, len(rows) + 10), cols=max(10, len(headers)))
            ws.resize(rows=max(100, len(rows) + 10), cols=max(10, len(headers)))
            ws.update(values=[headers] + [[str(c) if c is not None else "" for c in r] for r in rows], range_name="A1")
        return jsonify({"success": True, "message": f"ส่งออก {len(jobs)} ใบงานไป Google Sheet แล้ว (ชีต EXPORT_...)"})
    except Exception as e:
        return jsonify({"success": False, "message": f"ส่งออกไม่สำเร็จ: {e}"}), 500


# ====================== กันเซิร์ฟเวอร์หลับ (Render Free) ======================
def _keep_alive_loop(url):
    while True:
        time.sleep(10 * 60)
        try:
            requests.get(url, timeout=15)
        except Exception as e:
            print(f"⚠️ keep-alive ping ล้มเหลว: {e}")


def _start_background():
    load_settings()
    restore_state()
    try:
        _seen_groups.update(storage.get("seen_groups") or {})
    except Exception as e:
        print(f"⚠️ โหลดรายชื่อกลุ่มไม่สำเร็จ: {e}")
    try:
        persist_state()  # เขียนใบงานที่กู้คืนมาลงตาราง jobs ทันที (ย้ายข้อมูลเก่าเข้าตารางใหม่)
    except Exception as e:
        print(f"⚠️ บันทึกสถานะตอนเริ่มระบบไม่สำเร็จ: {e}")
    base = os.getenv("KEEP_ALIVE_URL") or os.getenv("RENDER_EXTERNAL_URL")
    if base and os.getenv("KEEP_ALIVE", "1") != "0":
        url = base.rstrip("/") + "/health"
        threading.Thread(target=_keep_alive_loop, args=(url,), daemon=True).start()
        print(f"🔔 เปิด keep-alive ping ไปที่ {url} ทุก 10 นาที")
    for b in bots_status():
        print(f"🤖 {b['name']}: {'รับข้อความได้ (webhook + secret พร้อม)' if b['receiving'] else 'ยังไม่มี secret ใน LINE_BACKUP_SECRETS (ส่งได้อย่างเดียว ยังรับข้อความไม่ได้)'}")
    if len(_backup_secrets) > len(BACKUP_CONFIGURATIONS):
        print("⚠️ LINE_BACKUP_SECRETS มีมากกว่า LINE_BACKUP_TOKENS เช็คลำดับให้ตรงกัน")
    print(f"💾 ที่เก็บข้อมูล: {'Supabase' if storage.using_supabase() else 'ไฟล์ในเครื่อง (local_store.json)'}")


_start_background()


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", 5000)), debug=False)
